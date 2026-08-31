"""합성 Golden Sample fixture (Design §8.5) — PDF·PPTX 각 5페이지, 결정적 생성.

구성: 1 표지(전면 배경 이미지+제목) / 2 목차 / 3 차트+해설 / 4 표 / 5 결론.
로고 PNG 는 2~5 페이지 우상단 동일 위치 반복 (→ logo 에셋).
표지 배경은 전면 (→ cover 에셋).
팔레트: 남색 #1F3A5F(제목) · 주황 #E07A1F(강조) · 본문 #222222.
"""

from __future__ import annotations

import io

import fitz

NAVY = (0x1F, 0x3A, 0x5F)
ORANGE = (0xE0, 0x7A, 0x1F)
TEXT = (0x22, 0x22, 0x22)
PAGE_W, PAGE_H = 960.0, 540.0  # pt (13.333 x 7.5 inch)


def png(w: int, h: int, color: tuple[int, int, int]) -> bytes:
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, w, h), False)
    pix.clear_with(0)
    for x in range(w):
        for y in range(h):
            pix.set_pixel(x, y, color)
    return pix.tobytes("png")


LOGO_PNG = png(120, 40, ORANGE)
COVER_PNG = png(192, 108, NAVY)
CHART_PNG = png(400, 240, (90, 160, 90))


def _rgb(c: tuple[int, int, int]) -> tuple[float, float, float]:
    return tuple(v / 255 for v in c)  # type: ignore[return-value]


def _text(page: fitz.Page, xy, text, size, color=TEXT, bold=False) -> None:
    page.insert_text(
        xy, text, fontsize=size, fontname="hebo" if bold else "helv", color=_rgb(color)
    )


def _logo(page: fitz.Page) -> None:
    page.insert_image(fitz.Rect(820, 20, 940, 60), stream=LOGO_PNG)


def _table(page: fitz.Page, rect: fitz.Rect, rows: int = 4, cols: int = 3) -> None:
    shape = page.new_shape()
    rw, ch = rect.width / cols, rect.height / rows
    for r in range(rows + 1):
        y = rect.y0 + r * ch
        shape.draw_line((rect.x0, y), (rect.x1, y))
    for c in range(cols + 1):
        x = rect.x0 + c * rw
        shape.draw_line((x, rect.y0), (x, rect.y1))
    shape.finish(width=1)
    shape.commit()
    for r in range(rows):
        for c in range(cols):
            _text(page, (rect.x0 + c * rw + 4, rect.y0 + r * ch + 16), f"r{r}c{c}", 12)


def sample_pdf() -> bytes:
    doc = fitz.open()
    p1 = doc.new_page(width=PAGE_W, height=PAGE_H)
    p1.insert_image(fitz.Rect(0, 0, PAGE_W, PAGE_H), stream=COVER_PNG)
    _text(p1, (80, 250), "Quarterly Risk Report", 28, (255, 255, 255), bold=True)
    _text(p1, (80, 300), "Credit Review Dept", 14, (255, 255, 255))

    p2 = doc.new_page(width=PAGE_W, height=PAGE_H)
    _logo(p2)
    _text(p2, (80, 80), "Contents", 20, NAVY, bold=True)
    for i, line in enumerate(["1. Overview", "2. Delinquency", "3. Actions"]):
        _text(p2, (100, 140 + i * 30), line, 14)

    p3 = doc.new_page(width=PAGE_W, height=PAGE_H)
    _logo(p3)
    _text(p3, (80, 80), "Delinquency Trend", 20, NAVY, bold=True)
    p3.insert_image(fitz.Rect(80, 120, 520, 384), stream=CHART_PNG)
    _text(p3, (560, 160), "Rate rose 0.3p in 3Q", 14)
    _text(p3, (560, 190), "Driven by SME segment", 14, ORANGE)

    p4 = doc.new_page(width=PAGE_W, height=PAGE_H)
    _logo(p4)
    _text(p4, (80, 80), "Portfolio Summary", 20, NAVY, bold=True)
    _table(p4, fitz.Rect(80, 120, 880, 400))

    p5 = doc.new_page(width=PAGE_W, height=PAGE_H)
    _logo(p5)
    _text(p5, (80, 80), "Next Steps", 20, NAVY, bold=True)
    _text(p5, (100, 140), "Tighten SME underwriting", 14)
    _text(p5, (400, 500), "Confidential", 10, (120, 120, 120))
    data = doc.tobytes()
    doc.close()
    return data


def sample_pptx() -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, box, txt, size, color, bold=False, font=None):
        tb = slide.shapes.add_textbox(*box)
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = txt
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font or ("HeadFont" if bold else "BodyFont")
        run.font.color.rgb = RGBColor(*color)
        return tb

    def logo(slide):
        slide.shapes.add_picture(
            io.BytesIO(LOGO_PNG), Inches(11.4), Inches(0.3), Inches(1.6), Inches(0.55)
        )

    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_picture(
        io.BytesIO(COVER_PNG), 0, 0, prs.slide_width, prs.slide_height
    )
    text(
        s1,
        (Inches(1), Inches(3), Inches(10), Inches(1)),
        "Quarterly Risk Report",
        28,
        (255, 255, 255),
        True,
    )
    text(
        s1,
        (Inches(1), Inches(4.2), Inches(10), Inches(0.6)),
        "Credit Review Dept",
        14,
        (255, 255, 255),
    )

    s2 = prs.slides.add_slide(blank)
    logo(s2)
    text(
        s2, (Inches(1), Inches(0.8), Inches(8), Inches(0.8)), "Contents", 20, NAVY, True
    )
    text(
        s2,
        (Inches(1.2), Inches(2), Inches(8), Inches(2)),
        "1. Overview\n2. Delinquency\n3. Actions",
        14,
        TEXT,
    )

    s3 = prs.slides.add_slide(blank)
    logo(s3)
    text(
        s3,
        (Inches(1), Inches(0.8), Inches(8), Inches(0.8)),
        "Delinquency Trend",
        20,
        NAVY,
        True,
    )
    cd = CategoryChartData()
    cd.categories = ["1Q", "2Q", "3Q"]
    cd.add_series("Rate", (1.1, 1.2, 1.5))
    s3.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1.8),
        Inches(6),
        Inches(4.5),
        cd,
    )
    text(
        s3,
        (Inches(7.5), Inches(2.2), Inches(5), Inches(2)),
        "Rate rose 0.3p in 3Q",
        14,
        TEXT,
    )
    text(
        s3,
        (Inches(7.5), Inches(3), Inches(5), Inches(1)),
        "Driven by SME segment",
        14,
        ORANGE,
    )

    s4 = prs.slides.add_slide(blank)
    logo(s4)
    text(
        s4,
        (Inches(1), Inches(0.8), Inches(8), Inches(0.8)),
        "Portfolio Summary",
        20,
        NAVY,
        True,
    )
    tbl = s4.shapes.add_table(4, 3, Inches(1), Inches(1.8), Inches(11), Inches(4)).table
    for r in range(4):
        for c in range(3):
            tbl.cell(r, c).text = f"r{r}c{c}"

    s5 = prs.slides.add_slide(blank)
    logo(s5)
    text(
        s5,
        (Inches(1), Inches(0.8), Inches(8), Inches(0.8)),
        "Next Steps",
        20,
        NAVY,
        True,
    )
    text(
        s5,
        (Inches(1.2), Inches(2), Inches(8), Inches(1)),
        "Tighten SME underwriting",
        14,
        TEXT,
    )
    text(
        s5,
        (Inches(5), Inches(6.8), Inches(3), Inches(0.4)),
        "Confidential",
        10,
        (120, 120, 120),
    )

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
