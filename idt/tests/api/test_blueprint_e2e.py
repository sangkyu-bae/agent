"""Design §8.4 L3 — 합성 샘플 → 추출 → 저장(aiosqlite) → 워커 생성 → PPTX 재오픈.

UseCase 를 mock 하지 않는다: 실 추출기·정책·repository·렌더러·컴파일러 노드를 관통하고
비전/작성 LLM 만 fake (실 LLM 은 test_blueprint_smoke.py).
"""

from __future__ import annotations

import io
import os
import tempfile
from collections.abc import AsyncGenerator
from datetime import UTC, datetime
from unittest.mock import MagicMock

import pytest
from langchain_core.messages import AIMessage, HumanMessage
from pptx import Presentation
from sqlalchemy.ext.asyncio import (
    AsyncSession,
    async_sessionmaker,
    create_async_engine,
)
from src.api.blueprint_di import (
    build_presentation_generation_use_case,
    build_sample_extractor_registry,
)
from src.application.agent_builder.supervisor_hooks import DefaultHooks
from src.application.agent_builder.workflow_compiler import WorkflowCompiler
from src.application.blueprint.admin_use_case import BlueprintAdminUseCase
from src.application.blueprint.extraction_use_case import (
    BlueprintExtractionUseCase,
    VisionSession,
)
from src.domain.agent_attachment.value_objects import StoredAttachment
from src.domain.agent_builder.schemas import WorkerDefinition
from src.domain.blueprint.schemas import (
    NarrativeDraft,
    NarrativeSectionDraft,
    PagePatternDraft,
    SlideContentDraft,
    SlidePlanDraft,
    SlotDraft,
)
from src.domain.llm.interfaces import LLMFactoryInterface
from src.domain.multimodal.interfaces import DescribeOutcome
from src.domain.multimodal.value_objects import MultimodalSettings
from src.infrastructure.blueprint.fonts import FontCatalog
from src.infrastructure.blueprint.llm.synthesizer import NarrativeSynthesizer
from src.infrastructure.blueprint.models import (  # noqa: F401 — metadata 등록
    DocumentBlueprintAssetModel,
    DocumentBlueprintModel,
)
from src.infrastructure.blueprint.repository import (
    BlueprintRepository,
    SessionScopedBlueprintRepository,
)
from src.infrastructure.blueprint.vision.page_classifier import VisionPageClassifier
from src.infrastructure.persistence.models.base import Base

from tests.fixtures.blueprint_samples import sample_pdf

KINDS = ["cover", "toc", "chart_with_notes", "table", "closing"]


# ── fakes ─────────────────────────────────────────────────────────────────────


def _pattern_draft(kind: str) -> PagePatternDraft:
    slots = [
        SlotDraft(kind="title", x=0.08, y=0.1, w=0.8, h=0.12, role="제목", max_chars=40)
    ]
    if kind == "chart_with_notes":
        slots.append(
            SlotDraft(
                kind="chart",
                x=0.08,
                y=0.25,
                w=0.45,
                h=0.55,
                role="차트",
                max_chars=None,
            )
        )
        slots.append(
            SlotDraft(
                kind="bullets", x=0.58, y=0.3, w=0.35, h=0.4, role="해설", max_chars=200
            )
        )
    elif kind == "table":
        slots.append(
            SlotDraft(
                kind="table", x=0.08, y=0.25, w=0.84, h=0.55, role="표", max_chars=None
            )
        )
    elif kind != "cover":
        slots.append(
            SlotDraft(
                kind="bullets", x=0.1, y=0.3, w=0.8, h=0.5, role="본문", max_chars=300
            )
        )
    return PagePatternDraft(kind=kind, slots=slots, layout_notes="")


class FakeVisionAdapter:
    def build_image_block(self, image, options):
        return {"type": "image", "page": image.page}

    async def describe_with(self, messages, schema):
        if schema is NarrativeDraft:
            draft = NarrativeDraft(
                sections=[
                    NarrativeSectionDraft(role="표지", pattern_ids=["p1"], guidance=""),
                    NarrativeSectionDraft(
                        role="분석", pattern_ids=["p3", "p4"], guidance="추이·표"
                    ),
                    NarrativeSectionDraft(role="결론", pattern_ids=["p5"], guidance=""),
                ],
                tone="보고체",
                language="ko",
            )
        else:
            page = next(
                b["page"]
                for b in messages[-1].content
                if isinstance(b, dict) and b.get("type") == "image"
            )
            draft = _pattern_draft(KINDS[page - 1])
        return DescribeOutcome(
            draft=draft, degraded_output_mode=False, output_mode="strict", usage=None
        )


class FakeVisionProvider:
    async def resolve(self, request_id: str) -> VisionSession:
        settings = MultimodalSettings(
            id="s",
            enabled=True,
            vision_model_id="m1",
            max_images_per_doc=50,
            min_image_px=100,
            min_area_ratio=0.02,
            concurrency=2,
            timeout_sec=5,
            output_language="ko",
            detail_level="brief",
            updated_at=datetime.now(UTC),
        )
        adapter = FakeVisionAdapter()
        return VisionSession(
            settings, VisionPageClassifier(adapter), NarrativeSynthesizer(adapter)
        )


class FakeWorkerLLM:
    """워커 LLM 흉내 — strict 구조화 출력만 지원 (StructuredCaller 경로)."""

    def with_structured_output(self, schema, method=None, strict=None, **kw):
        outer = self

        class _S:
            async def ainvoke(self, messages, config=None):
                return {
                    "raw": None,
                    "parsed": outer._answer(schema, messages),
                    "parsing_error": None,
                }

        return _S()

    @staticmethod
    def _answer(schema, messages):
        if schema is SlidePlanDraft:
            return SlidePlanDraft(
                slides=[
                    {
                        "pattern_id": "p1",
                        "title": "3Q 리스크 보고",
                        "intent": "",
                        "data_hint": "",
                    },
                    {
                        "pattern_id": "p3",
                        "title": "연체율 추이",
                        "intent": "",
                        "data_hint": "연체율",
                    },
                    {
                        "pattern_id": "p4",
                        "title": "포트폴리오",
                        "intent": "",
                        "data_hint": "잔액",
                    },
                    {
                        "pattern_id": "p5",
                        "title": "향후 계획",
                        "intent": "",
                        "data_hint": "",
                    },
                ]
            )
        human = messages[-1].content
        slots = [
            {
                "slot_id": "title",
                "text": "제목",
                "bullets": None,
                "table": None,
                "chart": None,
            }
        ]
        if "slot_id=chart" in human:
            slots.append(
                {
                    "slot_id": "chart",
                    "text": None,
                    "bullets": None,
                    "table": None,
                    "chart": {
                        "type": "bar",
                        "categories": ["1Q", "2Q", "3Q"],
                        "series": [{"name": "연체율", "values": [1.1, 1.2, 1.5]}],
                        "unit": "%",
                    },
                }
            )
            slots.append(
                {
                    "slot_id": "bullets",
                    "text": None,
                    "bullets": ["3Q 0.3p 상승", "SME 주도"],
                    "table": None,
                    "chart": None,
                }
            )
        elif "slot_id=table" in human:
            slots.append(
                {
                    "slot_id": "table",
                    "text": None,
                    "bullets": None,
                    "chart": None,
                    "table": {
                        "header": ["구분", "잔액"],
                        "rows": [["기업", "10"], ["가계", "20"]],
                    },
                }
            )
        elif "slot_id=bullets" in human:
            slots.append(
                {
                    "slot_id": "bullets",
                    "text": None,
                    "bullets": ["심사 강화"],
                    "table": None,
                    "chart": None,
                }
            )
        return SlideContentDraft(slots=slots)


class MemStore:
    def __init__(self):
        self.files: dict[str, bytes] = {}

    def save(self, *, file_bytes, filename, attachment_type, owner_user_id):
        fid = f"f{len(self.files) + 1}"
        self.files[fid] = file_bytes
        return StoredAttachment(
            file_id=fid,
            type=attachment_type,
            filename=filename,
            size=len(file_bytes),
            owner_user_id=owner_user_id,
            file_path="x",
        )


@pytest.fixture
async def session_factory() -> AsyncGenerator:
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp.name}", future=True)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        yield factory
    finally:
        await engine.dispose()
        try:
            os.unlink(tmp.name)
        except OSError:
            pass


@pytest.mark.asyncio
async def test_l3_extract_save_generate_roundtrip(session_factory, tmp_path):
    logger = MagicMock()
    # 1) 추출 (실 PyMuPDF + fake 비전)
    extraction = BlueprintExtractionUseCase(
        extractors=build_sample_extractor_registry(),
        vision=FakeVisionProvider(),
        fonts=FontCatalog(font_dir=None, default="NanumGothic"),
        logger=logger,
    )
    outcome = await extraction.run(sample_pdf(), "golden.pdf", 60, "req")
    assert [p.kind.value for p in outcome.blueprint.patterns] == KINDS
    assert sorted(a.kind for a in outcome.blueprint.assets) == ["cover", "logo"]

    # 2) 관리자 저장 (실 repository, aiosqlite)
    async with session_factory() as session:
        async with session.begin():
            admin = BlueprintAdminUseCase(
                BlueprintRepository(session, logger), MagicMock(), logger
            )
            saved = await admin.create(
                outcome.blueprint, outcome.assets, created_by="admin"
            )
    blueprint_id = saved.id

    # 3) 워커 노드 → 생성 (실 렌더러, fake 워커 LLM, 런타임 세션 스코프 repo)
    store = MemStore()
    generator = build_presentation_generation_use_case(
        conversion_adapter=None,
        attachment_store=store,
        logger=logger,
        input_max_chars=2000,
    )
    llm_factory = MagicMock(spec=LLMFactoryInterface)
    compiler = WorkflowCompiler(
        tool_factory=MagicMock(),
        llm_factory=llm_factory,
        logger=logger,
        hooks=DefaultHooks(),
        presentation_generator=generator,
        blueprint_repository=SessionScopedBlueprintRepository(session_factory, logger),
    )
    worker = WorkerDefinition(
        tool_id="presentation_generator",
        worker_id="ppt",
        description="",
        sort_order=0,
        tool_config={
            "blueprint_id": blueprint_id,
            "output_format": "pptx",
            "mcp_pptx_to_pdf_tool_id": "",
            "max_slides": 10,
        },
    )
    node = compiler._create_presentation_generator_node(
        FakeWorkerLLM(), worker, auth_ctx=None, request_id="req"
    )
    out = await node(
        {
            "messages": [
                HumanMessage(content="3분기 리스크 PPT"),
                AIMessage(content="연체율 1Q 1.1 2Q 1.2 3Q 1.5", name="search"),
            ],
            "token_usage": 0,
            "token_limit": 8000,
        }
    )
    text = out["messages"][0].content
    assert "생성 완료" in text and "4장" in text and "차트 1개" in text

    # 4) PPTX 재오픈 검증 — 패턴 순서·팔레트·폰트 매핑·네이티브 차트·로고
    prs = Presentation(io.BytesIO(store.files["f1"]))
    assert len(prs.slides) == 4
    pics0 = [sh for sh in prs.slides[0].shapes if sh.shape_type == 13]
    assert pics0 and pics0[0].width == prs.slide_width  # 표지 전면 cover 에셋
    title_run = next(
        r
        for sh in prs.slides[0].shapes
        if sh.has_text_frame
        for p in sh.text_frame.paragraphs
        for r in p.runs
    )
    assert title_run.font.name == saved.font_mapping[saved.style.fonts["heading"]]
    assert title_run.font.size.pt == saved.style.sizes["h1"]
    charts = [
        sh
        for sh in prs.slides[1].shapes
        if getattr(sh, "has_chart", False) and sh.has_chart
    ]
    assert len(charts) == 1
    assert str(
        charts[0].chart.plots[0].series[0].format.fill.fore_color.rgb
    ) == saved.style.palette["accent1"].lstrip("#")
    tables = [
        sh
        for sh in prs.slides[2].shapes
        if getattr(sh, "has_table", False) and sh.has_table
    ]
    assert len(tables) == 1 and str(
        tables[0].table.cell(0, 0).fill.fore_color.rgb
    ) == saved.style.palette["primary"].lstrip("#")
    logos = [sh for sh in prs.slides[1].shapes if sh.shape_type == 13]
    assert len(logos) == 1  # 로고 에셋 재사용


class DegradedWorkerLLM(FakeWorkerLLM):
    """§8.4 #2: 계획에 미존재 pattern + 슬롯 길이 초과 + 차트 길이 불일치 → degraded."""

    @staticmethod
    def _answer(schema, messages):
        if schema is SlidePlanDraft:
            return SlidePlanDraft(
                slides=[
                    {
                        "pattern_id": "p1",
                        "title": "표지",
                        "intent": "",
                        "data_hint": "",
                    },
                    {
                        "pattern_id": "zzz",
                        "title": "없는 패턴",
                        "intent": "",
                        "data_hint": "",
                    },
                    {
                        "pattern_id": "p3",
                        "title": "차트",
                        "intent": "",
                        "data_hint": "",
                    },
                ]
            )
        human = messages[-1].content
        slots = [
            {
                "slot_id": "title",
                "text": "제목",
                "bullets": None,
                "table": None,
                "chart": None,
            }
        ]
        if "slot_id=chart" in human:
            slots.append(
                {
                    "slot_id": "chart",
                    "text": None,
                    "bullets": None,
                    "table": None,
                    "chart": {
                        "type": "bar",
                        "categories": ["1Q", "2Q"],
                        "series": [{"name": "r", "values": [1.0]}],
                        "unit": None,
                    },
                }
            )
            slots.append(
                {
                    "slot_id": "bullets",
                    "text": None,
                    "bullets": ["x" * 500],
                    "table": None,
                    "chart": None,
                }
            )
        return SlideContentDraft(slots=slots)


@pytest.mark.asyncio
async def test_l3_degraded_plan_and_slots_still_produce_deck(session_factory, tmp_path):
    """Design §8.4 #2 — 계획 1장 제외 + 슬롯 2개 제외(warnings) + 전체는 성공."""
    logger = MagicMock()
    extraction = BlueprintExtractionUseCase(
        extractors=build_sample_extractor_registry(),
        vision=FakeVisionProvider(),
        fonts=FontCatalog(font_dir=None, default="NanumGothic"),
        logger=logger,
    )
    outcome = await extraction.run(sample_pdf(), "golden.pdf", 60, "req")
    async with session_factory() as session:
        async with session.begin():
            admin = BlueprintAdminUseCase(
                BlueprintRepository(session, logger), MagicMock(), logger
            )
            saved = await admin.create(
                outcome.blueprint, outcome.assets, created_by="admin"
            )
    store = MemStore()
    generator = build_presentation_generation_use_case(
        conversion_adapter=None,
        attachment_store=store,
        logger=logger,
        input_max_chars=2000,
    )
    compiler = WorkflowCompiler(
        tool_factory=MagicMock(),
        llm_factory=MagicMock(spec=LLMFactoryInterface),
        logger=logger,
        hooks=DefaultHooks(),
        presentation_generator=generator,
        blueprint_repository=SessionScopedBlueprintRepository(session_factory, logger),
    )
    worker = WorkerDefinition(
        tool_id="presentation_generator",
        worker_id="ppt",
        description="",
        sort_order=0,
        tool_config={
            "blueprint_id": saved.id,
            "output_format": "pptx",
            "mcp_pptx_to_pdf_tool_id": "",
            "max_slides": 10,
        },
    )
    node = compiler._create_presentation_generator_node(
        DegradedWorkerLLM(), worker, auth_ctx=None, request_id="req"
    )
    out = await node(
        {
            "messages": [HumanMessage(content="PPT")],
            "token_usage": 0,
            "token_limit": 8000,
        }
    )
    text = out["messages"][0].content
    assert "생성 완료" in text and "2장" in text and "[주의]" in text
    assert "zzz" in text and "chart" in text and "bullets" in text
    prs = Presentation(io.BytesIO(store.files["f1"]))
    assert len(prs.slides) == 2
    assert not [
        sh
        for sh in prs.slides[1].shapes
        if getattr(sh, "has_chart", False) and sh.has_chart
    ]
