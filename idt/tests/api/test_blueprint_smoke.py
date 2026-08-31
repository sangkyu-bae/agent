"""실 LLM 스모크 — Plan FR-19 / Design §8.4 #3.

합성 Golden Sample → 실 비전 분류(gpt-4o) → 실 서사 종합 → 실 계획/작성 → PPTX 재오픈.
기본 pytest 에서는 제외(addopts -m 'not llm').
실행: `pytest -m llm tests/api/test_blueprint_smoke.py`
비용 가드: 페이지 5·max_slides 4.
"""

from __future__ import annotations

import io
import os
from datetime import UTC, datetime
from unittest.mock import MagicMock

import pytest
from dotenv import load_dotenv
from pptx import Presentation
from src.api.blueprint_di import (
    build_presentation_generation_use_case,
    build_sample_extractor_registry,
)
from src.application.blueprint.extraction_use_case import (
    BlueprintExtractionUseCase,
    VisionSession,
)
from src.domain.agent_attachment.value_objects import StoredAttachment
from src.domain.blueprint.tool_config import PresentationGeneratorToolConfig
from src.domain.blueprint.value_objects import PatternKind
from src.domain.llm_model.entity import LlmModel
from src.domain.multimodal.value_objects import MultimodalSettings
from src.infrastructure.blueprint.fonts import FontCatalog
from src.infrastructure.blueprint.llm.synthesizer import NarrativeSynthesizer
from src.infrastructure.blueprint.vision.page_classifier import VisionPageClassifier
from src.infrastructure.llm.llm_factory import LLMFactory
from src.infrastructure.multimodal.vision.openai_vision_adapter import (
    OpenAIVisionAdapter,
)

from tests.fixtures.blueprint_samples import sample_pdf

load_dotenv()
pytestmark = pytest.mark.llm


def _model(name: str) -> LlmModel:
    now = datetime.now(UTC)
    return LlmModel(
        id=f"smoke-{name}",
        provider="openai",
        model_name=name,
        display_name=name,
        description=None,
        api_key_env="OPENAI_API_KEY",
        max_tokens=None,
        is_active=True,
        is_default=False,
        created_at=now,
        updated_at=now,
        supports_vision=True,
    )


class _Provider:
    async def resolve(self, request_id: str) -> VisionSession:
        settings = MultimodalSettings(
            id="s",
            enabled=True,
            vision_model_id="smoke",
            max_images_per_doc=50,
            min_image_px=100,
            min_area_ratio=0.02,
            concurrency=3,
            timeout_sec=90,
            output_language="ko",
            detail_level="brief",
            updated_at=datetime.now(UTC),
        )
        adapter = OpenAIVisionAdapter(LLMFactory(), _model("gpt-4o"), MagicMock())
        return VisionSession(
            settings, VisionPageClassifier(adapter), NarrativeSynthesizer(adapter)
        )


class _Store:
    def __init__(self):
        self.files = {}

    def save(self, *, file_bytes, filename, attachment_type, owner_user_id):
        fid = f"f{len(self.files) + 1}"
        self.files[fid] = file_bytes
        return StoredAttachment(
            file_id=fid,
            type=attachment_type,
            filename=filename,
            size=len(file_bytes),
            owner_user_id=owner_user_id,
            file_path="x",
        )


@pytest.mark.asyncio
@pytest.mark.skipif(not os.environ.get("OPENAI_API_KEY"), reason="OPENAI_API_KEY 없음")
async def test_real_openai_extract_then_generate():
    logger = MagicMock()
    extraction = BlueprintExtractionUseCase(
        extractors=build_sample_extractor_registry(),
        vision=_Provider(),
        fonts=FontCatalog(font_dir=None, default="NanumGothic"),
        logger=logger,
    )
    out = await extraction.run(sample_pdf(), "golden.pdf", 5, "smoke")
    bp = out.blueprint
    kinds = [p.kind for p in bp.patterns]
    assert out.classification[1] <= 1, bp.warnings
    assert kinds[0] is PatternKind.COVER, kinds
    assert any(
        k in (PatternKind.CHART_WITH_NOTES, PatternKind.IMAGE_WITH_NOTES) for k in kinds
    ), kinds
    assert PatternKind.TABLE in kinds, kinds
    assert bp.narrative.sections and all(s.pattern_ids for s in bp.narrative.sections)
    assert bp.style.palette["primary"] == "#1F3A5F"

    store = _Store()
    generator = build_presentation_generation_use_case(
        conversion_adapter=None,
        attachment_store=store,
        logger=logger,
        input_max_chars=4000,
    )
    llm = LLMFactory().create(_model("gpt-4o-mini"), temperature=0.0)
    result = await generator.generate(
        llm=llm,
        blueprint=bp,
        assets=out.assets,
        tool_config=PresentationGeneratorToolConfig(blueprint_id=bp.id, max_slides=4),
        evidence_block=(
            "연체율: 1Q 1.1%, 2Q 1.2%, 3Q 1.5%. 기업 잔액 120억, 가계 잔액 80억."
        ),
        conversation_block="3분기 리스크 보고 PPT 를 4장 이내로 만들어줘",
        user_instruction="4장 이내, 차트 포함",
        owner_user_id="u",
        request_id="smoke",
    )
    assert 1 <= result.slide_count <= 4, result
    prs = Presentation(io.BytesIO(store.files[result.file_id]))
    assert len(prs.slides) == result.slide_count
    texts = [
        r.text
        for s in prs.slides
        for sh in s.shapes
        if sh.has_text_frame
        for p in sh.text_frame.paragraphs
        for r in p.runs
    ]
    assert any(t.strip() for t in texts)
    print(
        "\nSMOKE",
        {
            "kinds": [k.value for k in kinds],
            "slides": result.slide_count,
            "charts": result.chart_count,
            "warnings": result.warnings,
            "usage": result.usage,
        },
    )
