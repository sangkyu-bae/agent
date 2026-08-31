"""골든 샘플 회귀 — blueprint-style-fidelity Plan §4.1 SC-1~SC-8.

samples/golden_sample_report.pdf → (실제 PdfStyleExtractor + 정책) → (비전 분류는 DB
스냅샷 패턴으로 고정) → 실제 PptxSlideRenderer → python-pptx 로 재오픈해 검증한다.
비전 LLM 은 호출하지 않는다 — 비결정성 제거, CI 에서 실행 가능.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from src.domain.blueprint.font_normalization import normalize_blueprint_fonts
from src.domain.blueprint.serialization import blueprint_from_dict
from src.domain.blueprint.value_objects import (
    ChartSpec,
    PatternKind,
    SlideContent,
    SlidePlan,
    SlotContent,
    TableSpec,
)
from src.infrastructure.blueprint.renderer.pptx_renderer import PptxSlideRenderer

from tests.application.blueprint.test_extraction_use_case import GoldenAdapter, _uc

_ROOT = Path(__file__).parents[3]
_PDF = _ROOT / "samples" / "golden_sample_report.pdf"
_V1 = _ROOT / "tests" / "fixtures" / "blueprint" / "golden_v1.json"
_MALGUN = {"Malgun Gothic Bold", "Malgun Gothic Regular", "Malgun Gothic"}


def _sc(i, pid, *slots):
    return SlideContent(SlidePlan(i, pid, "t", "", ""), tuple(slots), ())


def _text(sid, text):
    return SlotContent(sid, text, None, None, None)


def _bullets(sid, *items):
    return SlotContent(sid, None, tuple(items), None, None)


def _slides():
    chart = ChartSpec("bar", ("1Q", "2Q", "3Q"), (("중소기업", (1.9, 2.1, 2.4)),), "%")
    table = TableSpec(("구분", "잔액"), (("가계", "12,480"), ("중소기업", "9,870")))
    return [
        _sc(
            1,
            "p1",
            _text("title", "2026년 3분기 여신 리스크 보고"),
            _text("text", "부제"),
        ),
        _sc(
            2, "p2", _text("title", "목차"), _bullets("bullets", "요약", "분석", "대응")
        ),
        _sc(
            3,
            "p3",
            _text("title", "1. 요약"),
            _text("text", "핵심"),
            _bullets("bullets", "a"),
        ),
        _sc(
            4,
            "p4",
            _text("title", "2. 추이"),
            SlotContent("chart", None, None, None, chart),
        ),
        _sc(
            5,
            "p5",
            _text("title", "3. 현황"),
            SlotContent("table", None, None, table, None),
        ),
        _sc(
            6,
            "p6",
            _text("title", "4. 대응"),
            _bullets("bullets", "강화", "확대", "관리"),
        ),
    ]


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    import asyncio

    async def go():
        out = await _uc(GoldenAdapter(), tmp_path=tmp_path_factory.mktemp("bp")).run(
            _PDF.read_bytes(), _PDF.name, 20, "golden-regression"
        )
        bp = out.blueprint
        data = PptxSlideRenderer().render(bp, _slides(), out.assets, bp.font_mapping)
        return bp, Presentation(io.BytesIO(data))

    return asyncio.run(go())


def _runs(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    yield sh, p, r


def _pics(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]


def _rects(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]


def test_sc1_all_runs_use_original_malgun_gothic(rendered):
    _, prs = rendered
    fonts = {r.font.name for s in prs.slides for _, _, r in _runs(s)}
    assert fonts and fonts <= _MALGUN, fonts


def test_sc2_logo_positions_cover_vs_body(rendered):
    _, prs = rendered
    sw, sh = prs.slide_width, prs.slide_height
    cover_logo = [p for p in _pics(prs.slides[0]) if p.width < sw * 0.5]
    assert len(cover_logo) == 1
    assert abs(cover_logo[0].left / sw - 0.06) <= 0.02
    for slide in list(prs.slides)[1:]:
        (logo,) = _pics(slide)
        assert abs(logo.left / sw - 0.83) <= 0.02 and abs(logo.top / sh - 0.04) <= 0.02


def test_sc3_exactly_one_page_number_bottom_right_per_body_slide(rendered):
    _, prs = rendered
    total = len(prs.slides)
    for n, slide in enumerate(list(prs.slides)[1:], start=2):
        hits = [(sh, r) for sh, _, r in _runs(slide) if r.text == f"{n} / {total}"]
        assert len(hits) == 1
        sh, _ = hits[0]
        assert (sh.left + sh.width) / prs.slide_width >= 0.85
        assert sh.top / prs.slide_height >= 0.9


def test_sc4_original_footer_text_with_caption_style(rendered):
    bp, prs = rendered
    for slide in list(prs.slides)[1:]:
        run = next(
            r
            for _, _, r in _runs(slide)
            if r.text == "여신심사부 · 대외비 · 2026년 3분기"
        )
        assert run.font.size.pt == bp.style.sizes["caption"]
        assert str(run.font.color.rgb) == "666666"


def test_sc5_footer_band_and_cards_rendered(rendered):
    _, prs = rendered
    for slide in list(prs.slides)[1:]:
        bands = [
            r
            for r in _rects(slide)
            if r.top / prs.slide_height >= 0.9 and r.width / prs.slide_width >= 0.99
        ]
        assert len(bands) == 1 and str(bands[0].fill.fore_color.rgb) == "F3F4F6"
    assert len(_rects(prs.slides[5])) >= 1 + 3  # 띠 + 카드 3 (+ 줄무늬)
    assert len(_rects(prs.slides[2])) >= 2
    assert len(_rects(prs.slides[3])) == 1  # 차트 슬라이드엔 띠만 (막대 아님)


def test_sc6_toc_numbered_16pt_and_title_24pt_primary(rendered):
    _, prs = rendered
    toc = prs.slides[1]
    items = [r for _, _, r in _runs(toc) if r.text[:2] in ("1.", "2.", "3.")]
    assert [r.text for r in items] == ["1. 요약", "2. 분석", "3. 대응"]
    assert {r.font.size.pt for r in items} == {16.0}
    title = next(r for _, _, r in _runs(toc) if r.text == "목차")
    assert title.font.size.pt == 24.0 and title.font.bold
    assert str(title.font.color.rgb) == "1F3A5F"


def test_sc7_palette(rendered):
    bp, _ = rendered
    assert bp.style.palette["primary"] == "#1F3A5F"
    assert bp.style.palette["accent1"] == "#E07A1F"


def test_sc8_v1_snapshot_renders_without_decorations():
    bp = blueprint_from_dict(json.loads(_V1.read_text(encoding="utf-8")))
    assert bp.schema_version == 1
    data = PptxSlideRenderer().render(bp, _slides(), {}, bp.font_mapping)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 6
    assert all(_rects(s) == [] for s in prs.slides)
    total = len(prs.slides)
    assert [r.text for _, _, r in _runs(prs.slides[1])].count(f"2 / {total}") == 1


# ── pptx-font-fidelity §5.4 산출물 체크리스트 ─────────────────────────────────

_SUBFAMILY_SUFFIXES = (" Regular", " Bold", " Light", " Medium", " SemiBold")


def _typefaces(run) -> dict:
    from pptx.oxml.ns import qn

    rpr = run._r.get_or_add_rPr()
    return {
        tag: (None if rpr.find(qn(tag)) is None else rpr.find(qn(tag)).get("typeface"))
        for tag in ("a:latin", "a:ea", "a:cs")
    }


def test_fr01_typefaces_carry_no_subfamily_suffix(rendered):
    bp, prs = rendered
    fonts = {r.font.name for s in prs.slides for _, _, r in _runs(s)}
    assert fonts == {"Malgun Gothic"}
    for name in bp.font_mapping.values():
        assert not name.endswith(_SUBFAMILY_SUFFIXES), name


def test_fr02_every_run_sets_east_asian_typeface(rendered):
    _, prs = rendered
    for slide in prs.slides:
        for _, _, run in _runs(slide):
            faces = _typefaces(run)
            assert faces["a:ea"] == faces["a:latin"] == faces["a:cs"] == "Malgun Gothic"


def test_fr02_table_cells_set_east_asian_typeface(rendered):
    _, prs = rendered
    table = next(
        sh.table
        for s in prs.slides
        for sh in s.shapes
        if getattr(sh, "has_table", False) and sh.has_table
    )
    run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    assert _typefaces(run)["a:ea"] == "Malgun Gothic"


def test_fr04_every_body_slide_has_a_header_title(rendered):
    """표지를 제외한 모든 슬라이드에 h2 크기 헤더가 있다."""
    bp, prs = rendered
    h2 = bp.style.sizes["h2"]
    for slide in list(prs.slides)[1:]:
        headers = [
            r
            for _, _, r in _runs(slide)
            if r.font.size and r.font.size.pt == h2 and r.font.bold
        ]
        assert headers, f"헤더 없음: {[r.text for _, _, r in _runs(slide)]}"


def test_fr03_heading_renders_above_bullets_with_h3_style(rendered):
    """소제목이 있는 bullets 슬롯은 h3·bold·primary 문단을 먼저 그린다."""
    bp, _ = rendered
    slides = _slides()
    target = slides[5]  # p6 — bullets 슬라이드
    slots = tuple(
        SlotContent(c.slot_id, c.text, c.bullets, c.table, c.chart, "핵심 관찰")
        if c.bullets
        else c
        for c in target.slots
    )
    slides[5] = SlideContent(target.plan, slots, ())
    data = PptxSlideRenderer().render(bp, slides, {}, bp.font_mapping)
    slide = Presentation(io.BytesIO(data)).slides[5]
    heading = next(r for _, _, r in _runs(slide) if r.text == "핵심 관찰")
    assert heading.font.bold and str(heading.font.color.rgb) == "1F3A5F"
    assert heading.font.size.pt == bp.style.size("h3")
    assert _typefaces(heading)["a:ea"] == "Malgun Gothic"


# ── blueprint-font-mapping-migration §8.4 — 저장된 오염 블루프린트 회귀 ─────


def _polluted_v1_dict() -> dict:
    """FR-01 이전 저장 형태 — 서브패밀리명 + 항등 매핑 (실사례 재현)."""
    data = json.loads(_V1.read_text(encoding="utf-8"))
    data["font_mapping"] = {
        "Malgun Gothic Bold": "Malgun Gothic Bold",
        "Malgun Gothic Regular": "Malgun Gothic Regular",
    }
    return data


def test_sc1_polluted_blueprint_renders_valid_family_names():
    """Plan SC-1 — 항등 매핑을 가진 오염 블루프린트도 유효한 패밀리명으로 렌더된다."""
    bp = normalize_blueprint_fonts(blueprint_from_dict(_polluted_v1_dict()))
    data = PptxSlideRenderer().render(bp, _slides(), {}, bp.font_mapping)
    prs = Presentation(io.BytesIO(data))

    seen = set()
    for slide in prs.slides:
        for _, _, run in _runs(slide):
            faces = _typefaces(run)
            assert set(faces.values()) == {"Malgun Gothic"}, faces
            seen.update(faces.values())
    assert seen == {"Malgun Gothic"}


def test_polluted_blueprint_render_matches_shape_count_of_raw_load():
    """정규화는 폰트 외 렌더 결과를 바꾸지 않는다."""
    raw = blueprint_from_dict(_polluted_v1_dict())
    normalized = normalize_blueprint_fonts(raw)

    def _shapes(bp):
        data = PptxSlideRenderer().render(bp, _slides(), {}, bp.font_mapping)
        prs = Presentation(io.BytesIO(data))
        return [len(s.shapes) for s in prs.slides]

    assert _shapes(normalized) == _shapes(raw)


def test_v1_snapshot_mapping_still_resolves_after_collapse():
    """설계 §13.3 — fonts 만 정규화하면 NanumGothic 으로 가던 매핑이 빗나간다.

    키·값 동시 정규화(DR-2)로 접힌 뒤에도 조회가 성립해야 한다.
    """
    raw = json.loads(_V1.read_text(encoding="utf-8"))
    bp = normalize_blueprint_fonts(blueprint_from_dict(raw))
    data = PptxSlideRenderer().render(bp, _slides(), {}, bp.font_mapping)
    prs = Presentation(io.BytesIO(data))

    fonts = {r.font.name for s in prs.slides for _, _, r in _runs(s)}
    assert fonts == {"NanumGothic"}, fonts


# ── blueprint-slot-box-snap §8.4·8.5 — 슬롯 좌표 실측 스냅 ───────────────────

_SNAP_TEXT_KINDS = {"title", "text", "bullets"}
_TOL_IN = 0.1  # SC-1 허용 오차 (inch)


def _pdf_spans(pno: int):
    """원본 PDF 페이지의 텍스트 span 을 상대 좌표로 반환."""
    import pymupdf

    doc = pymupdf.open(_PDF)
    page = doc[pno]
    w, h = page.rect.width, page.rect.height
    out = []
    for block in page.get_text("dict")["blocks"]:
        if block["type"] != 0:
            continue
        for line in block["lines"]:
            for s in line["spans"]:
                if s["text"].strip():
                    x0, y0, x1, y1 = s["bbox"]
                    out.append(
                        (s["text"], round(s["size"], 1),
                         x0 / w, y0 / h, x1 / w, y1 / h)
                    )
    doc.close()
    return out


def _measured_union(pno: int, box, caption: float = 10.0):
    """슬롯 박스 안에 중심이 들어가는 PDF span 들의 union (푸터 제외)."""
    hits = [
        s
        for s in _pdf_spans(pno)
        if not (s[3] >= 0.88 and s[1] <= caption + 1)
        and box.x <= (s[2] + s[4]) / 2 <= box.x + box.w
        and box.y <= (s[3] + s[5]) / 2 <= box.y + box.h
    ]
    if not hits:
        return None
    return (min(s[2] for s in hits), min(s[3] for s in hits))


def test_sc1_every_text_slot_matches_measured_span_origin(rendered):
    """SC-1 — 스냅된 슬롯 원점이 원본 PDF 텍스트 위치와 일치한다."""
    bp, _ = rendered
    checked = 0
    for pattern in bp.patterns:
        for slot in pattern.slots:
            if slot.kind.value not in _SNAP_TEXT_KINDS:
                continue
            measured = _measured_union(pattern.sample_page - 1, slot.box)
            assert measured is not None, f"{pattern.id}/{slot.id}: 매칭 span 없음"
            dx = abs(slot.box.x - measured[0]) * bp.style.slide_size[0]
            dy = abs(slot.box.y - measured[1]) * bp.style.slide_size[1]
            assert dx <= _TOL_IN and dy <= _TOL_IN, (
                f"{pattern.id}/{slot.id}: Δx={dx:.3f}in Δy={dy:.3f}in"
            )
            checked += 1
    assert checked == 15, checked  # p6 bullets 분할로 13 → 15 (C-1)


def _partial_overlaps(bp, pattern) -> list[str]:
    decos = (*bp.style.common_decorations, *pattern.decorations)
    out = []
    for slot in pattern.slots:
        if slot.kind.value not in _SNAP_TEXT_KINDS:
            continue
        b = slot.box
        for d in decos:
            k = d.box
            separated = (
                b.y + b.h <= k.y
                or b.y >= k.y + k.h
                or b.x + b.w <= k.x
                or b.x >= k.x + k.w
            )
            enclosed = (
                k.x <= b.x
                and k.y <= b.y
                and k.x + k.w >= b.x + b.w
                and k.y + k.h >= b.y + b.h
            )
            if not separated and not enclosed:
                out.append(f"{pattern.id}/{slot.id} × {d.id}")
    return out


def test_sc2_no_partial_overlap_in_any_pattern(rendered):
    """SC-2 — 전 패턴에서 장식-텍스트 슬롯 부분 겹침 0건.

    blueprint-slot-content-fill C-1(슬롯 분할)이 p6 예외를 제거했다.
    직전 사이클에서는 p6(카드 3장)만 3건 남아 있었다.
    """
    bp, _ = rendered
    offenders = [o for p in bp.patterns for o in _partial_overlaps(bp, p)]
    assert offenders == [], offenders


def test_reported_page3_overlap_is_resolved(rendered):
    """사용자 신고 원본 — 3페이지 하이라이트 박스와 본문의 0.15in 겹침 해소."""
    bp, _ = rendered
    p3 = next(p for p in bp.patterns if p.id == "p3")
    assert _partial_overlaps(bp, p3) == []


def test_multi_card_pattern_is_split_into_one_slot_per_card(rendered):
    """SC-3·SC-4 — 카드 3장이 슬롯 3개로 쪼개지고 겹침이 사라졌다.

    직전 사이클(blueprint-slot-box-snap)이 이월한 한계의 해소를 고정한다.
    """
    bp, _ = rendered
    p6 = next(p for p in bp.patterns if p.id == "p6")

    bullets = [s for s in p6.slots if s.kind.value == "bullets"]
    assert [s.id for s in bullets] == ["bullets", "bullets2", "bullets3"]
    assert _partial_overlaps(bp, p6) == []

    decos = (*bp.style.common_decorations, *p6.decorations)
    for slot in bullets:  # 각 슬롯이 장식 하나에 완전히 포함
        b = slot.box
        assert any(
            d.box.x <= b.x
            and d.box.y <= b.y
            and d.box.x + d.box.w >= b.x + b.w
            and d.box.y + d.box.h >= b.y + b.h
            for d in decos
        ), slot


def test_sc3_cover_title_x_is_snapped_from_vision_estimate(rendered):
    """SC-3 — 표지 제목 x 가 0.200(추정) → 0.062(실측)."""
    bp, _ = rendered
    cover = next(p for p in bp.patterns if p.kind is PatternKind.COVER)
    title = cover.slot("title")
    assert title is not None
    assert abs(title.box.x - 0.062) <= 0.005, title.box


def test_sc5_footer_spans_are_not_inside_any_text_slot(rendered):
    """SC-5 — 푸터 밴드 좌표가 어떤 텍스트 슬롯에도 포함되지 않는다."""
    bp, _ = rendered
    for pattern in bp.patterns:
        for slot in pattern.slots:
            if slot.kind.value not in _SNAP_TEXT_KINDS:
                continue
            assert slot.box.y < 0.88, f"{pattern.id}/{slot.id} 가 푸터 밴드에 걸침"


def test_snap_covers_every_pattern_kind(rendered):
    """패턴 종류 5종 전부에서 스냅이 성공했다 (설계 시나리오 28)."""
    bp, _ = rendered
    kinds = {
        p.kind.value
        for p in bp.patterns
        if any(s.kind.value in _SNAP_TEXT_KINDS for s in p.slots)
    }
    assert kinds >= {"cover", "toc", "section_lead", "chart_with_notes", "table"}


def test_rendered_textbox_position_matches_original_pdf(rendered):
    """SC-1 실질 검증 — 렌더된 도형 좌표가 원본 PDF 텍스트와 0.1in 이내."""
    bp, prs = rendered
    w_in, h_in = bp.style.slide_size
    emu_in = 914400
    for index, pattern in enumerate(bp.patterns):
        title = pattern.slot("title")
        if title is None:
            continue
        measured = _measured_union(pattern.sample_page - 1, title.box)
        assert measured is not None
        shapes = [
            sh
            for sh in prs.slides[index].shapes
            if sh.has_text_frame and sh.text_frame.text.strip()
        ]
        assert shapes, f"slide {index + 1}: 텍스트 도형 없음"
        top_shape = min(shapes, key=lambda sh: sh.top)
        dx = abs(top_shape.left / emu_in - measured[0] * w_in)
        dy = abs(top_shape.top / emu_in - measured[1] * h_in)
        assert dx <= _TOL_IN and dy <= _TOL_IN, (
            f"slide {index + 1}: Δx={dx:.3f}in Δy={dy:.3f}in"
        )


def test_snap_does_not_change_shape_count(rendered):
    """좌표 외 렌더 결과는 변하지 않는다 (설계 시나리오 30) — 실측 기준선."""
    _, prs = rendered
    assert [len(s.shapes) for s in prs.slides] == [4, 6, 9, 6, 6, 12]


# ── blueprint-slot-content-fill §8.4·8.5 — 분할·목차·길이 폴백 ───────────────


def test_split_does_not_affect_other_patterns(rendered):
    """시나리오 38 — p1~p5 의 텍스트 슬롯 개수는 변하지 않는다 (과분할 0건)."""
    bp, _ = rendered
    counts = {
        p.id: sum(1 for s in p.slots if s.kind.value in _SNAP_TEXT_KINDS)
        for p in bp.patterns
    }
    assert counts == {"p1": 2, "p2": 2, "p3": 3, "p4": 2, "p5": 2, "p6": 4}


def test_toc_bullets_are_filled_from_other_slide_titles():
    """SC-1 / 시나리오 40 — 목차에 실제 슬라이드 제목이 채워진다."""
    from src.domain.blueprint.policies import TocContentPolicy

    bp = asyncio_run_extract()
    toc = next(p for p in bp.patterns if p.kind is PatternKind.TOC)
    plans = [
        SlidePlan(i, p.id, f"{i}. 섹션 {i}", "", "")
        for i, p in enumerate(bp.patterns, start=1)
    ]
    kinds = {p.id: p.kind for p in bp.patterns}
    toc_plan = next(pl for pl in plans if pl.pattern_id == toc.id)

    content, warnings = TocContentPolicy.apply(toc_plan, plans, kinds, toc)

    assert content is not None and warnings == ()
    bullets = next(c.bullets for c in content.slots if c.bullets)
    assert len(bullets) == 4  # 표지·목차 자신 제외한 4장
    # DR-6: 정책은 계획 제목을 **그대로** 담는다 — 접두 번호를 덧붙이지 않는다
    expected = tuple(
        pl.title for pl in plans if kinds[pl.pattern_id] not in
        (PatternKind.COVER, PatternKind.TOC)
    )
    assert bullets == expected


def asyncio_run_extract():
    import asyncio

    async def go():
        out = await _uc(GoldenAdapter()).run(
            _PDF.read_bytes(), _PDF.name, 20, "content-fill"
        )
        return out.blueprint

    return asyncio.run(go())


def test_toc_rendered_numbers_are_not_duplicated(rendered):
    """시나리오 41 (DR-6) — 렌더러가 번호를 붙이므로 '1. 1.' 이 없다."""
    _, prs = rendered
    for slide in prs.slides:
        for _, _, run in _runs(slide):
            assert not run.text.startswith(("1. 1.", "2. 2.", "3. 3.", "4. 4."))


def test_cards_each_have_their_own_slot(rendered):
    """SC-2 / 시나리오 42 — 카드 3장이 각자 슬롯을 갖는다 (빈 카드 0개 전제)."""
    bp, _ = rendered
    p6 = next(p for p in bp.patterns if p.id == "p6")
    bullets = [s for s in p6.slots if s.kind.value == "bullets"]

    assert len(bullets) == 3
    tops = [s.box.y for s in bullets]
    assert tops == sorted(tops) and len(set(tops)) == 3  # 서로 다른 위치


def test_over_max_chars_content_survives_to_render():
    """SC-5 / 시나리오 43 — max_chars 를 넘긴 내용이 렌더 결과에 살아남는다."""
    from src.domain.blueprint.policies import SlotContentPolicy

    bp = asyncio_run_extract()
    pattern = next(
        p for p in bp.patterns if any(s.kind.value == "bullets" for s in p.slots)
    )
    slot = next(s for s in pattern.slots if s.kind.value == "bullets")
    long_bullet = "가" * ((slot.max_chars or 10) + 50)

    outcome = SlotContentPolicy.apply(
        [SlotContent(slot.id, None, (long_bullet,), None, None)], pattern
    )
    assert len(outcome.kept) == 1 and outcome.warnings  # 유지 + 경고

    slides = [_sc(1, pattern.id, *outcome.kept)]
    data = PptxSlideRenderer().render(bp, slides, {}, bp.font_mapping)
    prs = Presentation(io.BytesIO(data))

    rendered = "".join(r.text for s in prs.slides for _, _, r in _runs(s))
    assert long_bullet in rendered


# ── blueprint-render-style-fidelity §8.4 시나리오 26~32 — 렌더 스타일 ───────

_GOLDEN_STYLE = dict(
    body_line_spacing=1.45,
    body_space_after_pt=33.2,
    chart_label_size_pt=9.0,
)
_GOLDEN_TABLE = dict(zebra=True, border_width_pt=0.75)


def _render_with_golden_style(bp):
    from dataclasses import replace

    style = replace(
        bp.style,
        table_style=replace(bp.style.table_style, **_GOLDEN_TABLE),
        **_GOLDEN_STYLE,
    )
    styled = replace(bp, style=style)
    data = PptxSlideRenderer().render(styled, _slides(), {}, styled.font_mapping)
    return styled, Presentation(io.BytesIO(data))


def test_golden_style_applies_chart_labels_and_unit(rendered):
    """SC-1·SC-2 / 시나리오 26~27."""
    bp, _ = rendered
    styled, prs = _render_with_golden_style(bp)

    charts = [
        sh.chart for s in prs.slides for sh in s.shapes if getattr(sh, "has_chart", 0)
    ]
    assert charts
    xml = charts[0]._chartSpace.xml
    assert "<c:dLbls" in xml
    assert '<c:numFmt formatCode="0.00&quot;%&quot;"' in xml  # _slides() 의 unit="%"


def test_golden_style_applies_table_borders_and_zebra(rendered):
    """SC-3·SC-4 / 시나리오 28~29."""
    from pptx.oxml.ns import qn

    bp, _ = rendered
    styled, prs = _render_with_golden_style(bp)

    tables = [
        sh.table for s in prs.slides for sh in s.shapes if getattr(sh, "has_table", 0)
    ]
    assert tables
    table = tables[0]
    tc_pr = table.cell(0, 0)._tc.find(qn("a:tcPr"))
    srgb = tc_pr.find(qn("a:lnL")).find(qn("a:solidFill")).find(qn("a:srgbClr"))
    assert srgb.get("val") == styled.style.table_style.border.lstrip("#").upper()
    assert str(table.cell(2, 0).fill.fore_color.rgb) == "F3F4F6"  # 데이터 2행


def test_golden_style_applies_paragraph_spacing(rendered):
    """SC-5 / 시나리오 30."""
    bp, _ = rendered
    _, prs = _render_with_golden_style(bp)

    spaced = [
        p
        for s in prs.slides
        for sh in s.shapes
        if sh.has_text_frame
        for p in sh.text_frame.paragraphs
        if p.runs and p.runs[0].text.startswith("•")
    ]
    assert spaced
    assert all(p.line_spacing == 1.45 for p in spaced)
    assert all(round(p.space_after.pt, 1) == 33.2 for p in spaced)


def test_style_tokens_do_not_change_shape_counts(rendered):
    """시나리오 31 — 스타일만 바뀌고 도형 구성은 동일하다.

    에셋을 양쪽 모두 비워 **스타일 차이만** 비교한다.
    """
    bp, _ = rendered
    base_data = PptxSlideRenderer().render(bp, _slides(), {}, bp.font_mapping)
    base = Presentation(io.BytesIO(base_data))
    _, styled = _render_with_golden_style(bp)

    assert [len(s.shapes) for s in styled.slides] == [
        len(s.shapes) for s in base.slides
    ]


def test_v1_snapshot_renders_with_default_style(rendered):
    """SC-6 / 시나리오 32 — v1 은 신규 필드가 없어도 예외 없이 렌더된다."""
    raw = json.loads(_V1.read_text(encoding="utf-8"))
    bp = normalize_blueprint_fonts(blueprint_from_dict(raw))

    assert bp.style.chart_label_size_pt == 0.0
    assert bp.style.table_style.border_width_pt == 0.0

    data = PptxSlideRenderer().render(bp, _slides(), {}, bp.font_mapping)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 6
