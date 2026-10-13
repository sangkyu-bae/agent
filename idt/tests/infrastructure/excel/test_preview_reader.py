"""ExcelPreviewReader 테스트 (excel-preview Design §4.3, §8.2 #2~5).

픽스처는 테스트 내 openpyxl 생성 (바이너리 커밋 금지 — Design §8.5).
"""
import datetime
import time

import pytest
from openpyxl import Workbook

from src.domain.excel_preview.exceptions import PreviewParseFailedError
from src.infrastructure.excel.preview_reader import ExcelPreviewReader

MAX_ROWS = 100
MAX_COLS = 50


@pytest.fixture
def reader() -> ExcelPreviewReader:
    return ExcelPreviewReader()


@pytest.fixture
def multi_sheet_path(tmp_path):
    """시트 3개: 일반 / 빈 시트 / 날짜·숫자 혼합."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "일반"
    ws1.append(["지점", "매출", "비고"])
    ws1.append(["서울", 1200, None])
    ws1.append(["부산", 900.5, "비수기"])

    wb.create_sheet("빈시트")

    ws3 = wb.create_sheet("혼합")
    ws3.append(["날짜", "정수형실수", "불리언"])
    ws3.append([datetime.datetime(2026, 10, 1, 9, 0, 0), 1.0, True])

    path = tmp_path / "multi.xlsx"
    wb.save(path)
    return str(path)


class TestReadMultiSheet:
    def test_sheet_order_and_names(self, reader, multi_sheet_path):
        sheets = reader.read(multi_sheet_path, MAX_ROWS, MAX_COLS)
        assert [s.name for s in sheets] == ["일반", "빈시트", "혼합"]

    def test_header_and_rows_serialized(self, reader, multi_sheet_path):
        s = reader.read(multi_sheet_path, MAX_ROWS, MAX_COLS)[0]
        assert s.columns == ["지점", "매출", "비고"]
        assert s.rows == [["서울", "1200", ""], ["부산", "900.5", "비수기"]]
        assert s.total_rows == 3
        assert s.total_cols == 3
        assert s.truncated is False

    def test_empty_sheet(self, reader, multi_sheet_path):
        s = reader.read(multi_sheet_path, MAX_ROWS, MAX_COLS)[1]
        assert s.columns == []
        assert s.rows == []
        assert s.total_rows == 0
        assert s.total_cols == 0
        assert s.truncated is False

    def test_datetime_float_bool_cells(self, reader, multi_sheet_path):
        s = reader.read(multi_sheet_path, MAX_ROWS, MAX_COLS)[2]
        assert s.rows == [["2026-10-01T09:00:00", "1", "true"]]


def _write_only_xlsx(path, n_rows: int, n_cols: int) -> None:
    """write_only 저장 — 빠르지만 dimensions 메타가 기록되지 않는다."""
    from openpyxl import Workbook as WB

    wb = WB(write_only=True)
    ws = wb.create_sheet("big")
    ws.append([f"c{i}" for i in range(n_cols)])
    for r in range(n_rows):
        ws.append([r] * n_cols)  # 숫자 셀 — shared string 비대 방지
    wb.save(path)


def _inject_dimension(path, ref: str) -> None:
    """sheet1.xml에 dimension 메타 주입 — 산출 노드(pandas 일반 모드)와 동일 구조.

    dimensions 없는 파일은 openpyxl load_workbook이 전체 시트 XML을 스캔하므로
    (openpyxl 3.1 _get_size → parse_dimensions), 대용량 픽스처는 메타가 있어야
    프로덕션과 같은 성능 특성을 가진다.
    """
    import os
    import zipfile

    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(
        tmp, "w", zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(
                    b"<sheetData>",
                    f'<dimension ref="{ref}"/><sheetData>'.encode(),
                    1,
                )
            zout.writestr(item, data)
    os.replace(tmp, path)


@pytest.fixture(scope="module")
def large_path(tmp_path_factory):
    """2만 행 × 60열, dimensions 메타 포함 (산출 노드 동일 조건). 모듈 스코프 1회."""
    path = tmp_path_factory.mktemp("large") / "large.xlsx"
    _write_only_xlsx(path, n_rows=20000, n_cols=60)
    _inject_dimension(path, "A1:BH20001")
    return str(path)


@pytest.fixture(scope="module")
def nodim_path(tmp_path_factory):
    """500행 × 60열, dimensions 메타 없음 — §3.4 폴백 케이스."""
    path = tmp_path_factory.mktemp("nodim") / "nodim.xlsx"
    _write_only_xlsx(path, n_rows=500, n_cols=60)
    return str(path)


class TestLimits:
    def test_row_and_col_limits(self, reader, large_path):
        s = reader.read(large_path, MAX_ROWS, MAX_COLS)[0]
        assert len(s.columns) == MAX_COLS
        assert len(s.rows) == MAX_ROWS
        assert all(len(r) == MAX_COLS for r in s.rows)
        assert s.total_rows == 20001  # 헤더 포함 실제 행 수 (메타)
        assert s.total_cols == 60
        assert s.truncated is True

    def test_no_dimensions_meta_falls_back_to_lower_bound(self, reader, nodim_path):
        """메타 없는 파일(write_only 산출 등): 읽은 행 수를 하한값으로 쓴다 (§3.4)."""
        s = reader.read(nodim_path, MAX_ROWS, MAX_COLS)[0]
        assert len(s.rows) == MAX_ROWS
        assert s.total_rows >= MAX_ROWS + 1  # 하한값 — 헤더 + 상한 + 중단 행
        assert s.truncated is True

    def test_read_elapsed_under_2s(self, reader, large_path):
        # Plan SC: 10MB 프리뷰 응답 2초 이내 — read_only 스트리밍이 상한에서 중단돼야 한다
        start = time.perf_counter()
        reader.read(large_path, MAX_ROWS, MAX_COLS)
        assert time.perf_counter() - start < 2.0

    def test_header_only_sheet_not_truncated(self, reader, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["h1", "h2"])
        path = tmp_path / "header_only.xlsx"
        wb.save(path)

        s = reader.read(str(path), MAX_ROWS, MAX_COLS)[0]
        assert s.columns == ["h1", "h2"]
        assert s.rows == []
        assert s.total_rows == 1
        assert s.truncated is False

    def test_rows_padded_to_header_width(self, reader, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["a", "b", "c"])
        ws.append(["x"])  # 짧은 행 → 헤더 폭으로 패딩
        path = tmp_path / "ragged.xlsx"
        wb.save(path)

        s = reader.read(str(path), MAX_ROWS, MAX_COLS)[0]
        assert s.rows == [["x", "", ""]]


class TestFormula:
    def test_formula_without_cache_serializes_empty(self, reader, tmp_path):
        """수식 캐시 없는 셀: data_only=True면 None → "" (Design §3.3, Plan 리스크 6)."""
        wb = Workbook()
        ws = wb.active
        ws.append(["계산"])
        ws.append(["=1+1"])  # openpyxl 직접 저장은 계산 캐시가 없다
        path = tmp_path / "formula.xlsx"
        wb.save(path)

        s = reader.read(str(path), MAX_ROWS, MAX_COLS)[0]
        assert s.rows == [[""]]


class TestCorrupt:
    def test_non_zip_bytes_raises_parse_failed(self, reader, tmp_path):
        path = tmp_path / "corrupt.xlsx"
        path.write_bytes(b"this is not a zip file")
        with pytest.raises(PreviewParseFailedError) as exc_info:
            reader.read(str(path), MAX_ROWS, MAX_COLS)
        assert exc_info.value.__cause__ is not None  # 원인 체이닝 보존
