"""PptxSlideRenderer (Design §9.3 / D1 / D6) — 생성 PPTX 재오픈 구조·스타일 검증."""

from __future__ import annotations

import io
from datetime import UTC, datetime

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu
from src.domain.blueprint.value_objects import (
    BlueprintAsset,
    ChartSpec,
    DocumentBlueprint,
    HeaderFooter,
    Narrative,
    NarrativeSection,
    PagePattern,
    PatternKind,
    RelBox,
    SlideContent,
    SlidePlan,
    Slot,
    SlotContent,
    SlotKind,
    StyleTokens,
    TableSpec,
    TableStyle,
)
from src.infrastructure.blueprint.renderer.pptx_renderer import PptxSlideRenderer

from tests.fixtures.blueprint_samples import COVER_PNG, LOGO_PNG


def _bp() -> DocumentBlueprint:
    now = datetime(2026, 8, 22, tzinfo=UTC)
    logo = BlueprintAsset(
        "logo",
        "logo",
        "image/png",
        120,
        40,
        "l" * 64,
        RelBox(0.85, 0.03, 0.12, 0.07),
        True,
    )
    cover = BlueprintAsset(
        "cover", "cover", "image/png", 192, 108, "c" * 64, RelBox(0, 0, 1, 1), True
    )
    patterns = (
        PagePattern(
            "cover",
            PatternKind.COVER,
            (
                Slot(
                    "title",
                    SlotKind.TITLE,
                    RelBox(0.08, 0.4, 0.84, 0.15),
                    "제목",
                    40,
                    None,
                    None,
                ),
                Slot(
                    "sub",
                    SlotKind.TEXT,
                    RelBox(0.08, 0.58, 0.84, 0.08),
                    "부제",
                    60,
                    None,
                    None,
                ),
            ),
            None,
            1,
            "",
        ),
        PagePattern(
            "chart",
            PatternKind.CHART_WITH_NOTES,
            (
                Slot(
                    "title",
                    SlotKind.TITLE,
                    RelBox(0.08, 0.08, 0.84, 0.12),
                    "제목",
                    40,
                    None,
                    None,
                ),
                Slot(
                    "chart",
                    SlotKind.CHART,
                    RelBox(0.08, 0.25, 0.5, 0.6),
                    "차트",
                    None,
                    None,
                    None,
                ),
                Slot(
                    "notes",
                    SlotKind.BULLETS,
                    RelBox(0.62, 0.25, 0.3, 0.6),
                    "해설",
                    300,
                    None,
                    None,
                ),
            ),
            None,
            3,
            "",
        ),
        PagePattern(
            "table",
            PatternKind.TABLE,
            (
                Slot(
                    "title",
                    SlotKind.TITLE,
                    RelBox(0.08, 0.08, 0.84, 0.12),
                    "제목",
                    40,
                    None,
                    None,
                ),
                Slot(
                    "tbl",
                    SlotKind.TABLE,
                    RelBox(0.08, 0.25, 0.84, 0.6),
                    "표",
                    None,
                    10,
                    None,
                ),
                Slot(
                    "img",
                    SlotKind.IMAGE,
                    RelBox(0.4, 0.9, 0.2, 0.08),
                    "로고",
                    None,
                    None,
                    "logo",
                ),
            ),
            "#F3F4F6",
            4,
            "",
        ),
    )
    return DocumentBlueprint(
        id="b",
        name="n",
        description="",
        schema_version=1,
        source_kind="pdf",
        page_count=5,
        style=StyleTokens(
            (13.333, 7.5),
            {"heading": "HY헤드라인M", "body": "맑은 고딕"},
            {"h1": 30.0, "h2": 22.0, "body": 14.0, "caption": 10.0},
            {
                "primary": "#1F3A5F",
                "accent1": "#E07A1F",
                "accent2": "#2A9D8F",
                "text": "#222222",
                "bg": "#FFFFFF",
            },
            TableStyle("#1F3A5F", "#FFFFFF", "#CCCCCC", True),
            HeaderFooter("logo", "{n} / {total}", "Confidential"),
        ),
        patterns=patterns,
        narrative=Narrative((NarrativeSection("표지", ("cover",), ""),), "", "ko"),
        assets=(logo, cover),
        font_mapping={"HY헤드라인M": "NanumGothicBold", "맑은 고딕": "NanumGothic"},
        warnings=(),
        status="active",
        created_at=now,
        updated_at=now,
    )


def _slides() -> list[SlideContent]:
    return [
        SlideContent(
            SlidePlan(1, "cover", "3Q 리스크 보고", "", ""),
            (
                SlotContent("title", "3Q 리스크 보고", None, None, None),
                SlotContent("sub", "여신심사부", None, None, None),
            ),
            (),
        ),
        SlideContent(
            SlidePlan(2, "chart", "연체율 추이", "", ""),
            (
                SlotContent("title", "연체율 추이", None, None, None),
                SlotContent(
                    "chart",
                    None,
                    None,
                    None,
                    ChartSpec(
                        "bar",
                        ("1Q", "2Q", "3Q"),
                        (("연체율", (1.1, 1.2, 1.5)), ("목표", (1.0, 1.0, 1.0))),
                        "%",
                    ),
                ),
                SlotContent(
                    "notes", None, ("3Q 0.3p 상승", "SME 부문 주도"), None, None
                ),
            ),
            (),
        ),
        SlideContent(
            SlidePlan(3, "table", "포트폴리오", "", ""),
            (
                SlotContent("title", "포트폴리오", None, None, None),
                SlotContent(
                    "tbl",
                    None,
                    None,
                    TableSpec(("구분", "잔액"), (("기업", "10"), ("가계", "20"))),
                    None,
                ),
            ),
            (),
        ),
    ]


def _render() -> Presentation:
    data = PptxSlideRenderer().render(
        _bp(), _slides(), {"logo": LOGO_PNG, "cover": COVER_PNG}, _bp().font_mapping
    )
    assert data[:2] == b"PK"
    return Presentation(io.BytesIO(data))


def _runs(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    yield sh, r


def test_slide_count_size_and_order():
    prs = _render()
    assert len(prs.slides) == 3
    assert round(prs.slide_width / 914400, 3) == 13.333 and prs.slide_height == Emu(
        int(7.5 * 914400)
    )
    titles = [next(r.text for _, r in _runs(s)) for s in prs.slides]
    assert titles == ["3Q 리스크 보고", "연체율 추이", "포트폴리오"]


def test_cover_has_full_background_picture_and_h1_mapped_font():
    s = _render().slides[0]
    pics = [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert pics and pics[0].left == 0 and pics[0].width == Emu(int(13.333 * 914400))
    _, title = next(_runs(s))
    assert (
        title.font.name == "NanumGothicBold"
        and title.font.size.pt == 30
        and title.font.bold
    )
    assert str(title.font.color.rgb) == "FFFFFF"  # 표지(배경 이미지) 위 제목은 흰색


def test_chart_slide_native_chart_with_palette_and_bullets_and_logo():
    s = _render().slides[1]
    charts = [sh for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart]
    assert len(charts) == 1
    chart = charts[0].chart
    assert chart.chart_type.name.startswith("COLUMN")
    series = list(chart.plots[0].series)
    assert [se.name for se in series] == ["연체율", "목표"]
    assert str(series[0].format.fill.fore_color.rgb) == "E07A1F"
    assert str(series[1].format.fill.fore_color.rgb) == "2A9D8F"
    texts = [r.text for _, r in _runs(s)]
    assert any(t.startswith("• 3Q 0.3p 상승") for t in texts)
    body = next(r for _, r in _runs(s) if r.text.startswith("• "))
    assert body.font.name == "NanumGothic" and body.font.size.pt == 14
    pics = [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 1  # 로고 (표지 아님)
    assert "2 / 3" in texts and "Confidential" in texts


def test_table_slide_header_style_background_and_image_slot():
    s = _render().slides[2]
    tables = [sh for sh in s.shapes if getattr(sh, "has_table", False) and sh.has_table]
    assert len(tables) == 1
    tbl = tables[0].table
    assert (
        len(tbl.rows) == 3
        and tbl.cell(0, 0).text == "구분"
        and tbl.cell(2, 1).text == "20"
    )
    assert str(tbl.cell(0, 0).fill.fore_color.rgb) == "1F3A5F"
    assert str(s.background.fill.fore_color.rgb) == "F3F4F6"
    pics = [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 2  # 헤더 로고 + image 슬롯 로고


def test_missing_slot_content_and_missing_pattern_are_skipped_not_fatal():
    bp = _bp()
    slides = [
        SlideContent(
            SlidePlan(1, "chart", "t", "", ""),
            (SlotContent("title", "only title", None, None, None),),
            (),
        ),
        SlideContent(SlidePlan(2, "nope", "t", "", ""), (), ()),
    ]
    prs = Presentation(io.BytesIO(PptxSlideRenderer().render(bp, slides, {}, {})))
    assert len(prs.slides) == 1  # 없는 패턴은 건너뜀
    assert not [
        sh
        for sh in prs.slides[0].shapes
        if getattr(sh, "has_chart", False) and sh.has_chart
    ]


def test_chart_count_helper():
    assert PptxSlideRenderer.count_charts(_slides()) == 1


# ── blueprint-style-fidelity §2.1 렌더 z-order / 푸터 단일화 / align / 역할별 크기 ──

from dataclasses import replace  # noqa: E402

from pptx.enum.dml import MSO_FILL_TYPE  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from src.domain.blueprint.value_objects import Decoration  # noqa: E402

_BAND = Decoration("common1", "rect", RelBox(0.0, 0.93, 1.0, 0.07), "#F3F4F6")
_CARD = Decoration(
    "deco1", "rect", RelBox(0.06, 0.22, 0.88, 0.15), "#F3F4F6", "#CCCCCC"
)


def _bp_v2() -> DocumentBlueprint:
    bp = _bp()
    style = replace(
        bp.style,
        sizes={**bp.style.sizes, "h3": 16.0, "subtitle": 18.0},
        header_footer=HeaderFooter(
            "logo",
            "{n} / {total}",
            "여신심사부 · 대외비",
            footer_box=RelBox(0.06, 0.945, 0.5, 0.04),
            page_number_box=RelBox(0.9, 0.945, 0.08, 0.04),
            footer_color="#666666",
        ),
        common_decorations=(_BAND,),
    )
    cover, chart, table = bp.patterns
    cover = replace(
        cover,
        slots=tuple(
            replace(s, align="center") if s.id == "title" else s for s in cover.slots
        ),
    )
    chart = replace(chart, decorations=(_CARD,))
    toc = PagePattern(
        "toc",
        PatternKind.TOC,
        (
            Slot(
                "title",
                SlotKind.TITLE,
                RelBox(0.05, 0.05, 0.9, 0.1),
                "제목",
                10,
                None,
                None,
            ),
            Slot(
                "items",
                SlotKind.BULLETS,
                RelBox(0.1, 0.2, 0.8, 0.6),
                "목차",
                100,
                None,
                None,
            ),
            Slot(
                "footer",
                SlotKind.FOOTER,
                RelBox(0.05, 0.9, 0.9, 0.05),
                "푸터",
                50,
                None,
                None,
            ),
        ),
        None,
        2,
        "",
    )
    logo = replace(bp.assets[0], cover_box=RelBox(0.06, 0.07, 0.19, 0.11))
    return replace(
        bp,
        schema_version=2,
        style=style,
        patterns=(cover, toc, chart, table),
        assets=(logo, bp.assets[1]),
    )


def _slides_v2() -> list[SlideContent]:
    cover, chart, table = _slides()
    toc = SlideContent(
        SlidePlan(2, "toc", "목차", "", ""),
        (
            SlotContent("title", "목차", None, None, None),
            SlotContent("items", None, ("요약", "분석", "대응"), None, None),
            SlotContent("footer", "LLM이 날조한 푸터", None, None, None),
        ),
        (),
    )
    return [
        cover,
        toc,
        replace(chart, plan=SlidePlan(3, "chart", "차트", "", "")),
        table,
    ]


def _render_v2() -> Presentation:
    bp = _bp_v2()
    data = PptxSlideRenderer().render(
        bp, _slides_v2(), {"logo": LOGO_PNG, "cover": COVER_PNG}, bp.font_mapping
    )
    return Presentation(io.BytesIO(data))


def _texts(slide):
    return [r.text for _, r in _runs(slide)]


def _rects(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]


def test_common_decoration_drawn_on_body_slides_before_content_not_on_cover():
    prs = _render_v2()
    assert _rects(prs.slides[0]) == []  # 표지 제외
    toc = prs.slides[1]
    rects = _rects(toc)
    assert len(rects) == 1
    band = rects[0]
    assert list(toc.shapes).index(band) == 0  # z-order: 장식이 콘텐츠보다 먼저
    assert str(band.fill.fore_color.rgb) == "F3F4F6"
    assert round(band.top / prs.slide_height, 2) == 0.93
    assert band.line.fill.type == MSO_FILL_TYPE.BACKGROUND  # 테두리 없음


def test_pattern_decoration_with_line_rendered_after_common():
    prs = _render_v2()
    chart = prs.slides[2]
    rects = _rects(chart)
    assert len(rects) == 2
    card = rects[1]
    assert round(card.top / prs.slide_height, 2) == 0.22
    assert str(card.line.color.rgb) == "CCCCCC"


def test_footer_rendered_once_from_style_and_llm_footer_slot_ignored():
    prs = _render_v2()
    toc = prs.slides[1]
    texts = _texts(toc)
    assert texts.count("2 / 4") == 1
    assert "여신심사부 · 대외비" in texts
    assert "LLM이 날조한 푸터" not in texts
    number = next(sh for sh, r in _runs(toc) if r.text == "2 / 4")
    assert number.left / prs.slide_width > 0.85 and number.top / prs.slide_height > 0.9
    run = next(r for _, r in _runs(toc) if r.text == "여신심사부 · 대외비")
    assert str(run.font.color.rgb) == "666666" and run.font.size.pt == 10.0


def test_logo_uses_cover_box_on_cover_and_body_box_elsewhere():
    prs = _render_v2()
    cover_pics = [
        sh for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    logo_on_cover = [
        p for p in cover_pics if round(p.width / prs.slide_width, 2) == 0.19
    ]
    assert (
        len(logo_on_cover) == 1
        and round(logo_on_cover[0].left / prs.slide_width, 2) == 0.06
    )
    body_pics = [
        sh for sh in prs.slides[1].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert round(body_pics[0].left / prs.slide_width, 2) == 0.85


def test_toc_bullets_numbered_with_h3_size_and_spacing():
    prs = _render_v2()
    toc = prs.slides[1]
    items = [r for _, r in _runs(toc) if r.text.startswith(("1.", "2.", "3."))]
    assert [r.text for r in items] == ["1. 요약", "2. 분석", "3. 대응"]
    assert all(r.font.size.pt == 16.0 for r in items)
    shape = next(sh for sh, r in _runs(toc) if r.text == "1. 요약")
    assert shape.text_frame.paragraphs[0].line_spacing == 1.5


def test_cover_title_centered_and_subtitle_uses_subtitle_size():
    prs = _render_v2()
    cover = prs.slides[0]
    title = next(sh for sh, r in _runs(cover) if r.text == "3Q 리스크 보고")
    assert title.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    sub = next(r for _, r in _runs(cover) if r.text == "여신심사부")
    assert sub.font.size.pt == 18.0


def test_v1_blueprint_without_v2_fields_still_renders_with_fallback_footer():
    prs = _render()  # _bp() 는 v1 (장식·box 없음)
    s = prs.slides[1]
    assert _rects(s) == []
    texts = _texts(s)
    assert texts.count("2 / 3") == 1 and "Confidential" in texts


def test_footer_boxes_widened_from_tight_span_bbox_keeping_anchor_edges():
    """추출된 푸터 box 는 글자 bbox(좁음) — 렌더 시 앵커 모서리를 지키며 넓힌다."""
    bp = _bp_v2()
    hf = replace(
        bp.style.header_footer,
        footer_box=RelBox(0.06, 0.95, 0.17, 0.02),
        page_number_box=RelBox(0.92, 0.95, 0.02, 0.02),
    )
    bp = replace(bp, style=replace(bp.style, header_footer=hf))
    prs = Presentation(
        io.BytesIO(
            PptxSlideRenderer().render(
                bp, _slides_v2(), {"logo": LOGO_PNG, "cover": COVER_PNG}, {}
            )
        )
    )
    toc = prs.slides[1]
    sw, sh_ = prs.slide_width, prs.slide_height
    number = next(sh for sh, r in _runs(toc) if r.text == "2 / 4")
    assert round((number.left + number.width) / sw, 2) == 0.94  # 오른쪽 모서리 유지
    assert round(number.width / sw, 2) >= 0.1 and not number.text_frame.word_wrap
    footer = next(sh for sh, r in _runs(toc) if r.text == "여신심사부 · 대외비")
    assert round(footer.left / sw, 2) == 0.06 and round(footer.width / sw, 2) >= 0.5
    assert round(footer.height / sh_, 2) >= 0.04


# ── pptx-font-fidelity FR-02/03/04 ────────────────────────────────────────────


def _rpr_typefaces(run):
    from pptx.oxml.ns import qn

    rpr = run._r.get_or_add_rPr()
    return {
        tag: (None if rpr.find(qn(tag)) is None else rpr.find(qn(tag)).get("typeface"))
        for tag in ("a:latin", "a:ea", "a:cs")
    }


def test_every_textbox_run_sets_latin_ea_and_cs_typefaces():
    """FR-02: 한글 글리프가 따르는 a:ea 가 latin 과 같은 폰트여야 한다."""
    prs = _render()
    for slide in prs.slides:
        for _, run in _runs(slide):
            faces = _rpr_typefaces(run)
            assert faces["a:latin"], f"latin 미설정: {run.text!r}"
            assert faces["a:ea"] == faces["a:latin"] == faces["a:cs"]


def test_table_cell_runs_also_set_east_asian_typeface():
    """FR-02: 표 셀도 같은 경로를 타야 한다."""
    slide = _render().slides[2]
    table = next(sh for sh in slide.shapes if getattr(sh, "has_table", False)).table
    run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    faces = _rpr_typefaces(run)
    assert faces["a:ea"] == faces["a:latin"] == "NanumGothic"


def _slides_with_heading(heading: str | None):
    return [
        SlideContent(
            SlidePlan(1, "chart", "연체율 추이", "", ""),
            (
                SlotContent("title", "연체율 추이", None, None, None),
                SlotContent(
                    "notes", None, ("3Q 0.3p 상승", "SME 주도"), None, None, heading
                ),
            ),
            (),
        )
    ]


def _render_slides(slides):
    bp = _bp()
    data = PptxSlideRenderer().render(
        bp, slides, {"logo": LOGO_PNG, "cover": COVER_PNG}, bp.font_mapping
    )
    return Presentation(io.BytesIO(data))


def test_bullets_heading_is_rendered_as_h3_subheading():
    """FR-03: 소제목은 h3 크기·heading 폰트·bold·primary 색."""
    slide = _render_slides(_slides_with_heading("핵심 관찰")).slides[0]
    heading = next(r for _, r in _runs(slide) if r.text == "핵심 관찰")
    assert heading.font.bold is True
    assert heading.font.name == "NanumGothicBold"
    assert str(heading.font.color.rgb) == "1F3A5F"
    assert heading.font.size.pt > 14.0  # body(14) 보다 크다
    bullet = next(r for _, r in _runs(slide) if r.text.startswith("• 3Q"))
    assert bullet.font.bold is False and bullet.font.size.pt == 14.0


def test_heading_and_bullets_live_in_one_textbox_in_order():
    slide = _render_slides(_slides_with_heading("핵심 관찰")).slides[0]
    box = next(sh for sh, r in _runs(slide) if r.text == "핵심 관찰")
    texts = [r.text for p in box.text_frame.paragraphs for r in p.runs]
    assert texts == ["핵심 관찰", "• 3Q 0.3p 상승", "• SME 주도"]


def test_bullets_without_heading_are_unchanged():
    slide = _render_slides(_slides_with_heading(None)).slides[0]
    texts = [r.text for _, r in _runs(slide)]
    assert "핵심 관찰" not in texts
    assert any(t.startswith("• 3Q") for t in texts)


def test_toc_pattern_ignores_heading():
    """TOC 는 번호 목록 규칙을 유지한다 (기존 동작 보존)."""
    bp = _bp_v2()
    slides = [
        SlideContent(
            SlidePlan(1, "toc", "목차", "", ""),
            (
                SlotContent("title", "목차", None, None, None),
                SlotContent(
                    "items", None, ("요약", "분석"), None, None, "무시될 소제목"
                ),
            ),
            (),
        )
    ]
    data = PptxSlideRenderer().render(bp, slides, {"logo": LOGO_PNG}, {})
    texts = [r.text for _, r in _runs(Presentation(io.BytesIO(data)).slides[0])]
    assert "무시될 소제목" not in texts
    assert "1. 요약" in texts and "2. 분석" in texts


def test_missing_title_content_falls_back_to_plan_title():
    """FR-04: title 슬롯 내용이 없으면 계획 제목으로 헤더를 그린다."""
    bp = _bp()
    slides = [
        SlideContent(
            SlidePlan(1, "chart", "2. 연체율 추이 분석", "", ""),
            (SlotContent("notes", None, ("a",), None, None),),
            (),
        )
    ]
    prs = Presentation(io.BytesIO(PptxSlideRenderer().render(bp, slides, {}, {})))
    texts = [r.text for _, r in _runs(prs.slides[0])]
    assert "2. 연체율 추이 분석" in texts


def test_cover_title_is_not_fabricated_from_plan_when_absent():
    """표지는 기존 동작 유지 — 슬롯 내용이 없으면 제목을 만들지 않는다."""
    bp = _bp()
    slides = [SlideContent(SlidePlan(1, "cover", "표지 제목", "", ""), (), ())]
    prs = Presentation(io.BytesIO(PptxSlideRenderer().render(bp, slides, {}, {})))
    assert "표지 제목" not in [r.text for _, r in _runs(prs.slides[0])]
