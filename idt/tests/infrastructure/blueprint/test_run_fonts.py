"""set_run_font (pptx-font-fidelity §12.1 / FR-02) — latin·ea·cs 3슬롯 설정."""

from pptx import Presentation
from pptx.oxml.ns import qn
from src.infrastructure.blueprint.renderer.run_fonts import set_run_font

_SLOTS = ("a:latin", "a:ea", "a:cs")


def _run():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    paragraph = box.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "한글 텍스트"
    return run


def _typefaces(run) -> dict[str, str | None]:
    rpr = run._r.get_or_add_rPr()
    out = {}
    for tag in _SLOTS:
        el = rpr.find(qn(tag))
        out[tag] = None if el is None else el.get("typeface")
    return out


def test_sets_latin_ea_and_cs_to_same_typeface():
    run = _run()
    set_run_font(run, "맑은 고딕")
    assert _typefaces(run) == dict.fromkeys(_SLOTS, "맑은 고딕")
    assert run.font.name == "맑은 고딕"


def test_second_call_updates_without_duplicating_elements():
    run = _run()
    set_run_font(run, "맑은 고딕")
    set_run_font(run, "NanumGothic")
    rpr = run._r.get_or_add_rPr()
    for tag in _SLOTS:
        assert len(rpr.findall(qn(tag))) == 1
    assert _typefaces(run) == dict.fromkeys(_SLOTS, "NanumGothic")


def test_empty_name_is_ignored():
    run = _run()
    set_run_font(run, "")
    assert _typefaces(run) == dict.fromkeys(_SLOTS, None)
