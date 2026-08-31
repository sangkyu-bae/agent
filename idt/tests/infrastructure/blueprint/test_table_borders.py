"""table_borders — 표 셀 테두리 oxml 조작 (python-pptx 미지원).

Design Ref: blueprint-render-style-fidelity §8.3 시나리오 15~16 (FR-03, DR-3).
"""

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu
from src.infrastructure.blueprint.renderer.table_borders import apply_borders

_SIDES = ("a:lnL", "a:lnR", "a:lnT", "a:lnB")


def _table(rows: int = 3, cols: int = 2):
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_table(rows, cols, 0, 0, Emu(6000000), Emu(2000000))
    return shape.table


def _line_elements(cell):
    tc_pr = cell._tc.find(qn("a:tcPr"))
    if tc_pr is None:
        return {}
    return {s: tc_pr.find(qn(s)) for s in _SIDES}


def test_borders_are_applied_to_every_cell_on_all_four_sides():
    """시나리오 15 — 각 셀의 네 변에 선이 생기고 색이 일치한다."""
    table = _table()

    apply_borders(table, "#CCCCCC", 0.75)

    for row in table.rows:
        for cell in row.cells:
            lines = _line_elements(cell)
            assert all(lines.get(s) is not None for s in _SIDES), lines
            for side in _SIDES:
                fill = lines[side].find(qn("a:solidFill"))
                assert fill is not None
                srgb = fill.find(qn("a:srgbClr"))
                assert srgb is not None and srgb.get("val") == "CCCCCC"


def test_border_width_is_converted_to_emu():
    table = _table(rows=2, cols=1)

    apply_borders(table, "#1F3A5F", 1.5)

    line = _line_elements(table.cell(0, 0))["a:lnL"]
    assert line.get("w") == str(int(1.5 * 12700))  # 1pt = 12700 EMU


def test_zero_width_draws_nothing():
    """시나리오 16 (FR-07) — 0 이면 현행 동작(테두리 없음)."""
    table = _table()

    apply_borders(table, "#CCCCCC", 0.0)

    for row in table.rows:
        for cell in row.cells:
            assert all(v is None for v in _line_elements(cell).values())


def test_negative_width_draws_nothing():
    table = _table()

    apply_borders(table, "#CCCCCC", -1.0)

    assert all(
        v is None for row in table.rows for cell in row.cells
        for v in _line_elements(cell).values()
    )


def test_reapplying_does_not_duplicate_elements():
    """멱등 — 두 번 적용해도 변 당 요소가 하나다."""
    table = _table(rows=2, cols=1)

    apply_borders(table, "#CCCCCC", 0.75)
    apply_borders(table, "#1F3A5F", 1.0)

    tc_pr = table.cell(0, 0)._tc.find(qn("a:tcPr"))
    for side in _SIDES:
        assert len(tc_pr.findall(qn(side))) == 1
    srgb = tc_pr.find(qn("a:lnL")).find(qn("a:solidFill")).find(qn("a:srgbClr"))
    assert srgb.get("val") == "1F3A5F"  # 마지막 값으로 갱신
