"""chart_builder — 데이터 레이블(D-1)·단위 서식(D-2).

Design Ref: blueprint-render-style-fidelity §8.3 시나리오 8~14 (FR-01·02·08).
"""

import io

import pytest
from pptx import Presentation
from pptx.util import Emu
from src.domain.blueprint.value_objects import (
    ChartSpec,
    HeaderFooter,
    StyleTokens,
    TableStyle,
)
from src.infrastructure.blueprint.renderer.chart_builder import add_chart


def _style(**overrides) -> StyleTokens:
    return StyleTokens(
        slide_size=(13.333, 7.5),
        fonts={"heading": "H", "body": "B"},
        sizes={"h1": 34.0, "h2": 24.0, "body": 13.0, "caption": 9.0},
        palette={
            "primary": "#1F3A5F",
            "accent1": "#E07A1F",
            "text": "#222222",
            "bg": "#FFFFFF",
        },
        table_style=TableStyle("#1F3A5F", "#FFFFFF", "#CCCCCC", False),
        header_footer=HeaderFooter(None, "{n}", ""),
        **overrides,
    )


def _render(spec: ChartSpec, style: StyleTokens) -> str:
    """차트를 렌더하고 저장 후 chartSpace XML 을 돌려준다."""
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_chart(slide, (0, 0, Emu(6000000), Emu(4000000)), spec, style)
    buf = io.BytesIO()
    prs.save(buf)
    reopened = Presentation(io.BytesIO(buf.getvalue()))
    return reopened.slides[0].shapes[0].chart._chartSpace.xml


def _bar(unit: str | None = None) -> ChartSpec:
    return ChartSpec("bar", ("1Q", "2Q"), (("중소기업", (1.62, 1.71)),), unit)


# ── 시나리오 8~9: 데이터 레이블 (FR-01) ────────────────────────────────────


def test_data_labels_are_added_when_label_size_is_set():
    xml = _render(_bar(), _style(chart_label_size_pt=9.0))

    assert "<c:dLbls" in xml


def test_no_data_labels_when_label_size_is_zero():
    """시나리오 9 (FR-07) — 기본값 0 이면 현행 동작."""
    xml = _render(_bar(), _style())

    assert "<c:dLbls" not in xml


def test_label_bold_flag_is_honoured():
    bold = _render(_bar(), _style(chart_label_size_pt=9.0, chart_label_bold=True))
    plain = _render(_bar(), _style(chart_label_size_pt=9.0, chart_label_bold=False))

    assert 'b="1"' in bold
    assert bold != plain


# ── 시나리오 10~13: 단위 서식 (FR-02, DR-4, FR-08) ─────────────────────────


def test_unit_is_applied_as_literal_suffix_not_percentage():
    """시나리오 10 (DR-4) — '0.00%' 는 백분율 변환이라 리터럴로 감싼다."""
    xml = _render(_bar("%"), _style(chart_label_size_pt=9.0))

    assert '<c:numFmt formatCode="0.00&quot;%&quot;"' in xml  # 리터럴
    assert 'formatCode="0.00%"' not in xml  # 백분율 변환 서식이 아니다


def test_unit_applies_to_value_axis_for_bar():
    xml = _render(_bar("%"), _style())

    assert "<c:valAx" in xml
    assert '<c:numFmt formatCode="0.0&quot;%&quot;"' in xml


def test_pie_chart_applies_unit_without_touching_value_axis():
    """시나리오 11 (DR-5) — pie 는 value_axis 접근이 ValueError 다."""
    pie = ChartSpec("pie", ("a", "b"), (("s", (1.0, 2.0)),), "%")

    xml = _render(pie, _style(chart_label_size_pt=9.0))  # 예외 없이 렌더

    assert "&quot;%&quot;" in xml


def test_unit_with_double_quote_does_not_break_format():
    """시나리오 12 (FR-08) — LLM 출력이 서식 문자열에 삽입된다."""
    xml = _render(_bar('크"기'), _style(chart_label_size_pt=9.0))

    assert "<c:dLbls" in xml  # 렌더 자체가 깨지지 않는다
    assert '크&quot;기' not in xml  # 큰따옴표가 서식에 그대로 새지 않는다


def test_no_unit_leaves_format_unset():
    xml = _render(_bar(None), _style(chart_label_size_pt=9.0))

    # python-pptx 가 캐시에 넣는 <c:formatCode>General</c:formatCode> 는 우리 것이 아님
    assert "<c:numFmt" not in xml


# ── 시나리오 14: 차트 종류별 레이블 위치 (DR-6) ────────────────────────────


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_every_chart_kind_renders_with_labels(kind):
    spec = ChartSpec(kind, ("a", "b"), (("s", (1.0, 2.0)),), "%")

    xml = _render(spec, _style(chart_label_size_pt=9.0))

    assert "<c:dLbls" in xml


def test_existing_behaviour_is_unchanged_with_default_style():
    """레이블·단위가 없으면 기존 렌더 결과와 동일한 요소 구성."""
    xml = _render(_bar(), _style())

    assert "<c:dLbls" not in xml and "<c:numFmt" not in xml


def test_unit_made_only_of_quotes_leaves_format_unset():
    """FR-08 방어 분기 — 따옴표를 제거하면 빈 문자열이 되는 경우."""
    xml = _render(_bar('""'), _style(chart_label_size_pt=9.0))

    assert "<c:numFmt" not in xml
