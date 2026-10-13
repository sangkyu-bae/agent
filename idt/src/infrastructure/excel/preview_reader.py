"""openpyxl 기반 .xlsx 프리뷰 리더 (excel-preview Design §4.3).

read_only + data_only 스트리밍 — 상위 max_rows행에서 중단해 대용량 파일에서도
빠르다 (Plan RISK: 대용량 파싱). total_rows/total_cols는 ws.max_row/max_column
메타 사용, 메타가 None인 비정형 파일만 읽은 행 수를 하한값으로 쓴다 (Design §3.4).
"""
from openpyxl import load_workbook

from src.domain.excel_preview.exceptions import PreviewParseFailedError
from src.domain.excel_preview.policies import serialize_cell
from src.domain.excel_preview.value_objects import ExcelSheetData


class ExcelPreviewReader:
    """동기 리더 — 호출자(유스케이스)가 asyncio.to_thread로 감싼다 (FR-07)."""

    def read(
        self, file_path: str, max_rows: int, max_cols: int
    ) -> list[ExcelSheetData]:
        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
        except Exception as e:
            raise PreviewParseFailedError(f"워크북을 열 수 없습니다: {e}") from e
        try:
            return [
                _read_sheet(ws, max_rows, max_cols) for ws in wb.worksheets
            ]
        except Exception as e:
            raise PreviewParseFailedError(f"시트를 읽을 수 없습니다: {e}") from e
        finally:
            wb.close()


class _SheetScan:
    """시트 1회 스트리밍 결과 — _read_sheet와 _stream_rows 사이 운반용."""

    __slots__ = ("header", "rows", "seen_rows", "max_raw_width", "hit_limit")

    def __init__(self) -> None:
        self.header: list[str] | None = None
        self.rows: list[list[str]] = []
        self.seen_rows = 0      # 실제로 만난 행 수 (헤더 포함)
        self.max_raw_width = 0  # 실제로 만난 최대 열 수 (메타 없음 폴백용)
        self.hit_limit = False


def _stream_rows(ws, max_col_arg: int | None, max_rows: int, max_cols: int) -> _SheetScan:
    """상한까지 행을 스트리밍 직렬화한다. 상한 도달 시 즉시 중단 (Plan RISK)."""
    scan = _SheetScan()
    for raw in ws.iter_rows(max_col=max_col_arg, values_only=True):
        scan.max_raw_width = max(scan.max_raw_width, len(raw))
        raw = raw[:max_cols]
        scan.seen_rows += 1
        if scan.header is None:
            scan.header = [serialize_cell(c) for c in raw]
            continue
        if len(scan.rows) >= max_rows:
            scan.hit_limit = True
            break
        scan.rows.append([serialize_cell(c) for c in raw])
    return scan


def _read_sheet(ws, max_rows: int, max_cols: int) -> ExcelSheetData:
    # 메타는 iter 전에 캡처 — dimensions 없는 파일(write_only 산출 등)은
    # 읽는 만큼 max_row가 갱신되는 lazy 값이라 iter 후엔 신뢰 불가 (Design §3.4)
    meta_rows = ws.max_row
    meta_cols = ws.max_column
    # max_col을 그대로 주면 openpyxl이 모든 행을 상한 폭까지 None 패딩한다
    # → 실제 폭(메타)과 상한의 min으로 요청하고, 메타 None이면 수동 슬라이스
    max_col_arg = min(max_cols, meta_cols) if meta_cols else None

    scan = _stream_rows(ws, max_col_arg, max_rows, max_cols)
    header = scan.header
    rows = scan.rows
    seen_rows, max_raw_width, hit_limit = (
        scan.seen_rows, scan.max_raw_width, scan.hit_limit
    )

    if header is None:  # 빈 시트 (FR-04)
        return ExcelSheetData(
            name=ws.title, columns=[], rows=[],
            total_rows=0, total_cols=0, truncated=False,
        )

    header, rows = _pad_to_grid(header, rows)
    total_rows = _resolve_total_rows(meta_rows, seen_rows, hit_limit)
    total_cols = meta_cols if meta_cols is not None else max_raw_width
    truncated = (
        hit_limit or total_rows - 1 > len(rows) or total_cols > max_cols
    )
    return ExcelSheetData(
        name=ws.title, columns=header, rows=rows,
        total_rows=total_rows, total_cols=total_cols, truncated=truncated,
    )


def _resolve_total_rows(
    meta_rows: int | None, seen_rows: int, hit_limit: bool
) -> int:
    """전체 행 수 산출 (Design §3.4).

    상한 전에 끝까지 읽었다면 seen_rows가 정확하다. 상한에서 중단했다면 메타를
    쓰되, 메타가 없거나 읽은 수보다 작으면(신뢰 불가) 읽은 수를 하한값으로 쓴다.
    """
    if not hit_limit:
        return seen_rows
    if meta_rows is not None and meta_rows >= seen_rows:
        return meta_rows
    return seen_rows


def _pad_to_grid(
    header: list[str], rows: list[list[str]]
) -> tuple[list[str], list[list[str]]]:
    """헤더·행을 동일 폭으로 패딩한다 — 그리드 렌더 일관성 (Design §3.1)."""
    width = max([len(header)] + [len(r) for r in rows])
    pad = lambda r: r + [""] * (width - len(r))  # noqa: E731
    return pad(header), [pad(r) for r in rows]
