"""python-pptx 도형 헬퍼 — 장식 사각형 (blueprint-style-fidelity §9.1 / DR-1).

렌더러 본체의 함수 길이 규칙(40줄)을 지키기 위해 분리. EMU 변환은 호출자 책임(D1).
"""

from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE

from src.infrastructure.blueprint.renderer.chart_builder import rgb


def add_rect(
    slide, emu: tuple[int, int, int, int], fill: str, line: str | None
) -> None:
    """채움 사각형. line 이 None 이면 테두리 없음."""
    x, y, w, h = emu
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
    shape.shadow.inherit = False
    if shape.has_text_frame:
        shape.text_frame.text = ""
