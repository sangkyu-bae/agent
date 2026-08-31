"""PptxSlideRenderer — blueprint + 슬라이드 내용 → PPTX 바이트 (python-pptx).

Design Ref: golden-sample-blueprint §9.3 SlideRendererPort / D1 (0..1 → EMU 변환은
여기만) / D6 (차트 3종+표) / §7 (LLM 출력은 데이터로만 렌더 — 마크업 해석 없음)
- 슬롯 종류별 렌더: title/text/bullets → 텍스트박스, table → 표,
  chart → 네이티브 차트, image → 에셋 그림. 패턴에 없는 슬롯·내용 없는 슬롯은
  건너뛴다(degraded 는 UseCase 가 기록).
- 없는 pattern_id 슬라이드는 건너뛴다(계획 검증은 UseCase 책임, 렌더러는 방어만).

blueprint-style-fidelity §2.1 (z-order): 배경 → 장식(common + 패턴) → 로고 → 슬롯
→ 푸터/페이지번호 1회. footer 슬롯은 LLM 내용을 무시하고 스타일 값으로만 그린다(DR-2).
v1 블루프린트(장식·box 없음)는 폴백 상수로 동일하게 렌더된다.
"""

from __future__ import annotations

import io
from collections.abc import Mapping, Sequence
from typing import NamedTuple

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from src.domain.blueprint.value_objects import (
    DocumentBlueprint,
    PagePattern,
    PatternKind,
    RelBox,
    SlideContent,
    Slot,
    SlotContent,
    SlotKind,
    StyleTokens,
)
from src.infrastructure.blueprint.renderer.chart_builder import add_chart, rgb
from src.infrastructure.blueprint.renderer.run_fonts import set_run_font
from src.infrastructure.blueprint.renderer.shapes import add_rect
from src.infrastructure.blueprint.renderer.table_borders import apply_borders

_EMU_PER_INCH = 914400
_WHITE = "#FFFFFF"
_BLANK_LAYOUT = 6
# v1 폴백 (HeaderFooter.*_box 가 None 일 때)
_PAGE_NUMBER_BOX = RelBox(0.86, 0.93, 0.12, 0.05)
_FOOTER_BOX = RelBox(0.04, 0.93, 0.5, 0.05)
_TOC_LINE_SPACING = 1.5
_MIN_NUMBER_W, _MIN_FOOTER_W, _MIN_FOOTER_H = 0.1, 0.5, 0.04
_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_HEADING_SPACE_AFTER = 6.0  # 소제목 아래 여백(pt) — 본문과 시각 분리 (FR-03)


class _Line(NamedTuple):
    """텍스트박스 한 문단의 내용과 스타일."""

    text: str
    font_role: str
    size_pt: float
    bold: bool
    color: str | None
    space_after_pt: float = 0.0


class PptxSlideRenderer:
    def render(
        self,
        blueprint: DocumentBlueprint,
        slides: Sequence[SlideContent],
        assets: Mapping[str, bytes],
        font_mapping: Mapping[str, str],
    ) -> bytes:
        prs = Presentation()
        w_in, h_in = blueprint.style.slide_size
        prs.slide_width = Emu(int(w_in * _EMU_PER_INCH))
        prs.slide_height = Emu(int(h_in * _EMU_PER_INCH))
        ctx = _Ctx(blueprint, assets, font_mapping, prs.slide_width, prs.slide_height)
        renderable = [s for s in slides if blueprint.pattern(s.plan.pattern_id)]
        for index, content in enumerate(renderable, start=1):
            pattern = blueprint.pattern(content.plan.pattern_id)
            assert pattern is not None
            slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])
            _render_slide(ctx, slide, pattern, content, index, len(renderable))
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    @staticmethod
    def count_charts(slides: Sequence[SlideContent]) -> int:
        return sum(1 for s in slides for c in s.slots if c.chart is not None)


class _Ctx:
    def __init__(
        self,
        blueprint: DocumentBlueprint,
        assets: Mapping[str, bytes],
        font_mapping: Mapping[str, str],
        width: int,
        height: int,
    ) -> None:
        self.bp = blueprint
        self.style: StyleTokens = blueprint.style
        self.assets = assets
        self.font_mapping = font_mapping
        self.width = width
        self.height = height
        self.adopted = {a.id: a for a in blueprint.adopted_assets()}

    def emu(self, box: RelBox) -> tuple[int, int, int, int]:
        return (
            int(box.x * self.width),
            int(box.y * self.height),
            int(box.w * self.width),
            int(box.h * self.height),
        )

    def font(self, role: str) -> str:
        original = self.style.fonts.get(role, "")
        return self.font_mapping.get(original, original) or original

    def asset_bytes(self, asset_id: str | None) -> bytes | None:
        if asset_id is None or asset_id not in self.adopted:
            return None
        return self.assets.get(asset_id)


# ── slide ─────────────────────────────────────────────────────────────────────


def _render_slide(
    ctx: _Ctx, slide, pattern: PagePattern, content: SlideContent, n: int, total: int
) -> None:
    is_cover = pattern.kind is PatternKind.COVER
    on_image = _background(ctx, slide, pattern, is_cover)
    _decorations(ctx, slide, pattern, is_cover)
    _logo(ctx, slide, is_cover)
    by_slot = {c.slot_id: c for c in content.slots}
    for slot in pattern.slots:
        if slot.kind is SlotKind.FOOTER:
            continue  # DR-2: 푸터는 스타일 값으로만 (LLM 내용 무시)
        _render_slot(
            ctx,
            slide,
            pattern,
            slot,
            _with_title_fallback(by_slot.get(slot.id), slot, content, is_cover),
            on_image,
        )
    if not is_cover:
        _footer(ctx, slide, n, total)


def _with_title_fallback(
    content: SlotContent | None, slot: Slot, slide_content: SlideContent, is_cover: bool
) -> SlotContent | None:
    """비표지 슬라이드의 섹션 헤더 누락 방지 — 계획 제목으로 채운다 (FR-04).

    표지는 기존 동작 유지: 제목을 만들어 내지 않는다.
    """
    if slot.kind is not SlotKind.TITLE or is_cover:
        return content
    if content is not None and content.text:
        return content
    title = slide_content.plan.title.strip()
    if not title:
        return content
    return SlotContent(slot.id, title, None, None, None)


def _background(ctx: _Ctx, slide, pattern: PagePattern, is_cover: bool) -> bool:
    """배경 처리. 표지 전면 이미지가 깔리면 True (제목을 흰색으로)."""
    if is_cover:
        cover = next((a for a in ctx.adopted.values() if a.kind == "cover"), None)
        data = ctx.asset_bytes(cover.id) if cover else None
        if data:
            slide.shapes.add_picture(io.BytesIO(data), 0, 0, ctx.width, ctx.height)
            return True
    color = pattern.background or (ctx.style.palette["primary"] if is_cover else None)
    if color:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color)
        return is_cover
    return False


def _decorations(ctx: _Ctx, slide, pattern: PagePattern, is_cover: bool) -> None:
    """공통 장식(비표지) → 패턴 장식 순. 콘텐츠보다 먼저 그려 z-order 하단에 둔다."""
    common = () if is_cover else ctx.style.common_decorations
    for deco in (*common, *pattern.decorations):
        add_rect(slide, ctx.emu(deco.box), deco.fill, deco.line)


def _logo(ctx: _Ctx, slide, is_cover: bool) -> None:
    logo_id = ctx.style.header_footer.logo_asset_id
    data = ctx.asset_bytes(logo_id)
    if not data or logo_id not in ctx.adopted:
        return
    asset = ctx.adopted[logo_id]
    if is_cover and asset.cover_box is None:
        return  # v1: 표지 위치 정보 없음 → 표지엔 그리지 않음 (현행 유지)
    box = asset.cover_box if is_cover else asset.box
    x, y, w, h = ctx.emu(box)  # type: ignore[arg-type]
    slide.shapes.add_picture(io.BytesIO(data), x, y, w, h)


def _footer(ctx: _Ctx, slide, n: int, total: int) -> None:
    """푸터/페이지번호 1회 (DR-2). 추출 box 는 글자 bbox 라 앵커를 지키며 넓힌다."""
    hf = ctx.style.header_footer
    size = ctx.style.sizes["caption"]
    color = hf.footer_color
    if hf.page_number_format:
        text = hf.page_number_format.replace("{n}", str(n)).replace(
            "{total}", str(total)
        )
        box = _widen(hf.page_number_box or _PAGE_NUMBER_BOX, _MIN_NUMBER_W, "right")
        _textbox(
            ctx, slide, box, [text], "body", size, False, color, "right", wrap=False
        )
    if hf.footer_text:
        box = _widen(hf.footer_box or _FOOTER_BOX, _MIN_FOOTER_W, "left")
        _textbox(
            ctx, slide, box, [hf.footer_text], "body", size, False, color, wrap=False
        )


def _widen(box: RelBox, min_w: float, anchor: str) -> RelBox:
    """폭·높이를 최소값까지 키운다 (anchor='right' 면 오른쪽 모서리 고정)."""
    w = max(box.w, min_w)
    h = max(box.h, _MIN_FOOTER_H)
    x = box.x + box.w - w if anchor == "right" else box.x
    y = box.y + box.h / 2 - h / 2
    x, y = min(max(x, 0.0), 1.0 - w), min(max(y, 0.0), 1.0 - h)
    return RelBox(x=x, y=y, w=w, h=h)


# ── slots ─────────────────────────────────────────────────────────────────────


def _render_slot(
    ctx: _Ctx,
    slide,
    pattern: PagePattern,
    slot: Slot,
    content: SlotContent | None,
    on_image: bool,
) -> None:
    if slot.kind is SlotKind.IMAGE:
        data = ctx.asset_bytes(slot.asset_id)
        if data:
            x, y, w, h = ctx.emu(slot.box)
            slide.shapes.add_picture(io.BytesIO(data), x, y, w, h)
        return
    if content is None:
        return
    if slot.kind is SlotKind.TITLE and content.text is not None:
        _title(ctx, slide, pattern, slot, content.text, on_image)
    elif slot.kind is SlotKind.TEXT and content.text is not None:
        _text(ctx, slide, pattern, slot, content.text, on_image)
    elif slot.kind is SlotKind.BULLETS and content.bullets:
        _bullets(ctx, slide, pattern, slot, content.bullets, content.heading)
    elif slot.kind is SlotKind.TABLE and content.table is not None:
        _table(ctx, slide, slot.box, content.table.header, content.table.rows)
    elif slot.kind is SlotKind.CHART and content.chart is not None:
        add_chart(slide, ctx.emu(slot.box), content.chart, ctx.style)


def _title(ctx, slide, pattern: PagePattern, slot: Slot, text: str, on_image) -> None:
    is_cover = pattern.kind is PatternKind.COVER
    size = ctx.style.size("h1" if is_cover else "h2")
    color = _WHITE if on_image else ctx.style.palette["primary"]
    _textbox(ctx, slide, slot.box, [text], "heading", size, True, color, slot.align)


def _text(ctx, slide, pattern: PagePattern, slot: Slot, text: str, on_image) -> None:
    """표지의 text 슬롯은 부제(subtitle 크기), 그 외는 본문."""
    role = "subtitle" if pattern.kind is PatternKind.COVER else "body"
    color = _WHITE if on_image else None
    _textbox(
        ctx,
        slide,
        slot.box,
        text.split("\n"),
        "body",
        ctx.style.size(role),
        False,
        color,
        slot.align,
    )


def _bullets(
    ctx, slide, pattern: PagePattern, slot: Slot, bullets, heading: str | None = None
) -> None:
    """toc 패턴: '1. ' 번호 + h3 크기 + 1.5 줄간격. 그 외: '• ' + body.

    toc 를 제외하면 heading 이 있을 때 소제목 문단을 맨 앞에 둔다 (FR-03).
    """
    if pattern.kind is PatternKind.TOC:
        lines = [
            _Line(f"{i}. {b}", "body", ctx.style.size("h3"), False, None)
            for i, b in enumerate(bullets, start=1)
        ]
        spacing = _TOC_LINE_SPACING
    else:
        body = ctx.style.size("body")
        # Design Ref: blueprint-render-style-fidelity D-4 / DR-7 — 비-TOC 불릿에만
        # 토큰 줄간격·문단 여백을 적용한다. 기본값(1.0 / 0)이면 현행과 동일하다.
        gap = ctx.style.body_space_after_pt
        lines = [_Line(f"• {b}", "body", body, False, None, gap) for b in bullets]
        spacing = ctx.style.body_line_spacing
        if heading:
            lines.insert(0, _heading_line(ctx, heading))
    _styled_textbox(ctx, slide, slot.box, lines, slot.align, spacing)


def _heading_line(ctx: _Ctx, heading: str) -> _Line:
    """소제목 문단 — h3 크기·heading 폰트·bold·primary (FR-03)."""
    return _Line(
        text=heading,
        font_role="heading",
        size_pt=ctx.style.size("h3"),
        bold=True,
        color=ctx.style.palette["primary"],
        space_after_pt=_HEADING_SPACE_AFTER,
    )


def _textbox(
    ctx: _Ctx,
    slide,
    box: RelBox,
    lines: Sequence[str],
    font_role: str,
    size_pt: float,
    bold: bool,
    color: str | None,
    align: str = "left",
    line_spacing: float = 1.0,
    wrap: bool = True,
) -> None:
    """단일 스타일 텍스트박스."""
    styled = [_Line(line, font_role, size_pt, bold, color) for line in lines]
    _styled_textbox(ctx, slide, box, styled, align, line_spacing, wrap)


def _styled_textbox(
    ctx: _Ctx,
    slide,
    box: RelBox,
    lines: Sequence[_Line],
    align: str = "left",
    line_spacing: float = 1.0,
    wrap: bool = True,
) -> None:
    """문단마다 스타일이 다를 수 있는 텍스트박스 (소제목 + 본문 혼합용)."""
    x, y, w, h = ctx.emu(box)
    tb = slide.shapes.add_textbox(x, y, w, h)
    frame = tb.text_frame
    frame.word_wrap = wrap
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = _ALIGN.get(align, PP_ALIGN.LEFT)
        if line_spacing != 1.0:
            paragraph.line_spacing = line_spacing
        if line.space_after_pt:
            paragraph.space_after = Pt(line.space_after_pt)
        run = paragraph.add_run()
        run.text = line.text
        set_run_font(run, ctx.font(line.font_role))
        run.font.size = Pt(line.size_pt)
        run.font.bold = line.bold
        run.font.color.rgb = rgb(line.color or ctx.style.palette["text"])


def _table(ctx: _Ctx, slide, box: RelBox, header, rows) -> None:
    x, y, w, h = ctx.emu(box)
    table = slide.shapes.add_table(len(rows) + 1, len(header), x, y, w, h).table
    ts = ctx.style.table_style
    for c, text in enumerate(header):
        cell = table.cell(0, c)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(ts.header_bg)
        _style_cell(ctx, cell, ts.header_text, bold=True)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            if ts.zebra and r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(ts.zebra_bg)  # D-3: 상수 제거
            _style_cell(ctx, cell, ctx.style.palette["text"], bold=False)
    # Design Ref: blueprint-render-style-fidelity D-3 / DR-3 — oxml 은 전용 모듈에서
    apply_borders(table, ts.border, ts.border_width_pt)


def _style_cell(ctx: _Ctx, cell, color: str, bold: bool) -> None:
    for paragraph in cell.text_frame.paragraphs:
        for run in paragraph.runs:
            set_run_font(run, ctx.font("body"))
            run.font.size = Pt(ctx.style.sizes["body"])
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
