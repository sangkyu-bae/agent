"""PptxStyleExtractor — python-pptx 로 Golden Sample PPTX 원본의 도형 통계를 뽑는다.

Design Ref: golden-sample-blueprint §2.3 / Plan FR-03 / D4
- 텍스트 run(폰트·크기·굵기·색) / 그림(blob sha) / 표·차트 영역 / 테마 major·minor 폰트.
- 렌더 PNG 는 없음(render_png=None) — 분류는 도형 통계 힌트 기반 (D4).
- 손상 파일 → SampleExtractionError.
"""

from __future__ import annotations

import hashlib
import io

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from src.domain.blueprint.errors import SampleExtractionError
from src.domain.blueprint.value_objects import (
    ImageRef,
    PageStats,
    RelBox,
    SampleStats,
    TextSpan,
)

_EMU_PER_INCH = 914400
_DEFAULT_SIZE_PT = 18.0
_DEFAULT_COLOR = "#000000"
_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


class PptxStyleExtractor:
    supported_extensions: tuple[str, ...] = ("pptx",)

    def extract(self, data: bytes, filename: str, max_pages: int) -> SampleStats:
        try:
            prs = Presentation(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001 — 손상 파일은 415 로 매핑
            raise SampleExtractionError(f"cannot open pptx: {exc}") from exc
        try:
            sw, sh = int(prs.slide_width), int(prs.slide_height)
            theme = _theme_fonts(prs)
            pages = tuple(
                _slide_stats(slide, i + 1, sw, sh, theme)
                for i, slide in enumerate(prs.slides)
                if i < max_pages
            )
            if not pages:
                raise SampleExtractionError("pptx has no slides")
            return SampleStats(
                source_kind="pptx",
                page_size=(sw / _EMU_PER_INCH, sh / _EMU_PER_INCH),
                pages=pages,
                theme_fonts=theme,
            )
        except SampleExtractionError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise SampleExtractionError(f"pptx parse failed: {exc}") from exc


# ── internal ──────────────────────────────────────────────────────────────────


def _rel(shape, sw: int, sh: int) -> RelBox | None:
    if shape.left is None or shape.width is None:
        return None
    x0 = max(0.0, int(shape.left) / sw)
    y0 = max(0.0, int(shape.top) / sh)
    x1 = min(1.0, (int(shape.left) + int(shape.width)) / sw)
    y1 = min(1.0, (int(shape.top) + int(shape.height)) / sh)
    if x1 - x0 <= 0 or y1 - y0 <= 0:
        return None
    return RelBox(x=x0, y=y0, w=x1 - x0, h=y1 - y0)


def _slide_stats(
    slide, number: int, sw: int, sh: int, theme: dict[str, str]
) -> PageStats:
    spans: list[TextSpan] = []
    images: list[ImageRef] = []
    tables: list[RelBox] = []
    charts: list[RelBox] = []
    for shape in slide.shapes:
        box = _rel(shape, sw, sh)
        if box is None:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            images.append(_image(shape, box))
        elif getattr(shape, "has_table", False) and shape.has_table:
            tables.append(box)
        elif getattr(shape, "has_chart", False) and shape.has_chart:
            charts.append(box)
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            spans.extend(_runs(shape, box, theme))
    return PageStats(
        number=number,
        width=sw / _EMU_PER_INCH * 72.0,
        height=sh / _EMU_PER_INCH * 72.0,
        spans=tuple(spans),
        images=tuple(images),
        tables=tuple(tables),
        has_text=any(s.text.strip() for s in spans),
        render_png=None,
        charts=tuple(charts),
    )


def _runs(shape, box: RelBox, theme: dict[str, str]):
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            text = run.text
            if not text.strip():
                continue
            size = run.font.size.pt if run.font.size is not None else _DEFAULT_SIZE_PT
            bold = bool(run.font.bold)
            yield TextSpan(
                text=text,
                font=run.font.name or theme.get("minor", ""),
                size=round(float(size), 1),
                bold=bold,
                color=_color(run),
                box=box,
            )


def _color(run) -> str:
    try:
        rgb = run.font.color.rgb
    except AttributeError:
        return _DEFAULT_COLOR
    return f"#{rgb}" if rgb is not None else _DEFAULT_COLOR


def _image(shape, box: RelBox) -> ImageRef:
    blob = shape.image.blob
    width, height = shape.image.size
    return ImageRef(
        sha256=hashlib.sha256(blob).hexdigest(),
        mime=shape.image.content_type,
        width=int(width),
        height=int(height),
        box=box,
        data=blob,
    )


def _theme_fonts(prs) -> dict[str, str]:
    """slide master 의 theme part 에서 major/minor latin typeface 를 읽는다.

    없으면 빈 문자열.
    """
    fonts = {"major": "", "minor": ""}
    try:
        master_part = prs.slide_master.part
        theme_part = next(
            (
                rel.target_part
                for rel in master_part.rels.values()
                if "theme" in rel.reltype
            ),
            None,
        )
        if theme_part is None:
            return fonts
        root = etree.fromstring(theme_part.blob)
        for key in ("major", "minor"):
            node = root.find(f".//a:{key}Font/a:latin", _NS)
            if node is not None:
                fonts[key] = node.get("typeface", "")
    except Exception:  # noqa: BLE001 — 테마는 선택 정보
        return fonts
    return fonts


__all__ = ["PptxStyleExtractor", "Emu"]
